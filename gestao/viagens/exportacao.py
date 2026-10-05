"""Planilha da lista de ofícios (paridade com o "Exportar" da referência).

O recorte é o mesmo da tela — busca, situação, filtros avançados e ordem —, então quem
exporta leva exatamente o que estava vendo. As 14 colunas seguem a referência; datas e
valores saem como data e número de verdade (a planilha soma e ordena).
"""

from __future__ import annotations

from collections.abc import Iterable
from io import BytesIO

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import services
from .models import Oficio
from .queries import trechos_de, viajantes_de

COLUNAS = ["Nº", "Data do ofício", "Protocolo", "Situação", "Tipo", "Destinos", "Saída",
           "Retorno", "Servidores", "Motorista", "Viatura", "Diárias (R$)",
           "Quantidade de diárias", "Justificativa"]
LARGURAS = [10, 14, 14, 12, 26, 40, 17, 17, 48, 28, 30, 14, 22, 14]
# Cinza neutro do cabeçalho: valor fixo do arquivo Excel, não cor de tela (tokens.css).
CABECALHO = PatternFill("solid", fgColor="D1D3D4")


def _destinos(oficio: Oficio, trechos) -> str:
    vistos: list[str] = []
    for t in trechos:
        rotulo = f"{t.destino.nome}/{t.destino.uf}"
        if t.destino_id != oficio.sede_id and rotulo not in vistos:
            vistos.append(rotulo)
    return ", ".join(vistos)


def _local(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def _transporte(oficio: Oficio) -> str:
    if oficio.viatura_id and oficio.viatura:
        return str(oficio.viatura)
    partes = [oficio.transporte_descricao, oficio.transporte_placa]
    return " · ".join(p for p in partes if p)


def _justificativa(oficio: Oficio) -> str:
    if (oficio.justificativa or "").strip():
        return "Preenchida"
    try:
        exigida = services.avaliar_prazo_do_oficio(oficio).justificativa_obrigatoria
    except services.RegraViolada:  # unidade sem configuração: sem regra de prazo
        exigida = False
    return "Pendente" if exigida else ""


def _texto_seguro(valor):
    """Texto que começa com = + - @ (ou tabulação/CR) seria fórmula no Excel: vai como
    texto literal (injeção de fórmula, revisão de segurança)."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


def linha(oficio: Oficio) -> list:
    trechos = trechos_de(oficio)
    equipe = viajantes_de(oficio)
    assunto = services.assunto_do_oficio(oficio)
    motorista = services.nome_do_motorista(oficio, equipe)
    return [_texto_seguro(v) for v in [
        oficio.numero_formatado,
        oficio.data_oficio,
        oficio.protocolo_formatado,
        oficio.get_situacao_display(),
        ("Autorização" if assunto.autorizacao else "Convalidação")
        + (f" {assunto.rotulo}" if oficio.marcador else ""),
        _destinos(oficio, trechos),
        _local(trechos[0].saida_em) if trechos else None,
        _local(trechos[-1].chegada_em) if trechos else None,
        ", ".join(v.servidor.nome for v in equipe),
        motorista,
        _transporte(oficio),
        oficio.diarias_total,
        oficio.diarias_resumo or "",
        _justificativa(oficio),
    ]]


def planilha_de_oficios(oficios: Iterable[Oficio]) -> bytes:
    livro = Workbook()
    aba = livro.active
    if aba is None:  # Workbook() sempre nasce com uma aba; isto só acalma o verificador
        aba = livro.create_sheet()
    aba.title = "Ofícios"
    aba.append(COLUNAS)
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = CABECALHO
    for oficio in oficios:
        aba.append(linha(oficio))
    for linha_planilha in aba.iter_rows(min_row=2):
        linha_planilha[1].number_format = "DD/MM/YYYY"
        linha_planilha[6].number_format = linha_planilha[7].number_format = "DD/MM/YYYY HH:MM"
        linha_planilha[11].number_format = "#,##0.00"
    for indice, largura in enumerate(LARGURAS, start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


COLUNAS_PRESTACAO = ["Servidor", "Ofício", "Protocolo", "Solicitação", "Liberação das diárias",
                     "Prazo limite de saque", "Prestar contas até", "Diária liberada",
                     "Situação", "O que falta"]


def planilha_de_prestacoes(linhas) -> bytes:
    """Prestações (referência: aba "Prestações", mesmas colunas)."""
    from . import diario, prestacao, relatorio
    from .dominio.prestacao import prazo_para_prestar

    linhas = list(linhas)
    ids = {ps.prestacao_id for ps in linhas}
    com_diario, com_rt = diario.preenchidos(ids), relatorio.preenchidos(ids)

    livro = Workbook()
    aba = livro.active or livro.create_sheet()
    aba.title = "Prestações"
    aba.append(COLUNAS_PRESTACAO)
    for celula in aba[1]:
        celula.font = Font(bold=True)
        celula.fill = CABECALHO
    for ps in linhas:
        oficio = ps.prestacao.oficio
        situacao = ("Finalizada" if ps.finalizada and ps.situacao in ("pendente", "preenchimento")
                    else "Arquivada" if ps.arquivada else ps.get_situacao_display())
        aba.append([_texto_seguro(v) for v in [
            ps.servidor.nome, oficio.numero_formatado, oficio.protocolo_formatado,
            ps.numero_solicitacao, ps.data_liberacao_diarias, ps.prazo_limite_saque,
            prazo_para_prestar(ps.prazo_limite_saque), prestacao.diaria_liberada(ps), situacao,
            "; ".join(prestacao.pendencias(ps, ps.prestacao_id in com_diario,
                                           ps.prestacao_id in com_rt))]])
    for linha_planilha in aba.iter_rows(min_row=2):
        for i in (4, 5, 6):
            linha_planilha[i].number_format = "DD/MM/YYYY"
        linha_planilha[7].number_format = "#,##0.00"
    for indice, largura in enumerate((32, 12, 16, 16, 14, 14, 14, 14, 16, 60), start=1):
        aba.column_dimensions[get_column_letter(indice)].width = largura
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
