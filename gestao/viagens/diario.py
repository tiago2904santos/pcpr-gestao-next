"""Diário de bordo (módulo 9b, paridade com `diario_services` da referência; ficha em
docs/migration/prestacao.md).

Um por prestação (a equipe compartilha). As linhas espelham os trechos do ofício e
guardam o que já foi digitado quando os trechos mudam. Motorista e viatura podem ser
trocados só aqui. Trava quando a equipe inteira está finalizada (o colega que ainda não
prestou contas precisa do diário).
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import F, Q
from django.db.models.functions import Coalesce
from django.utils import timezone

from . import policies
from .dominio import diario as dominio
from .models import (
    DiarioBordo,
    DiarioBordoTrecho,
    Oficio,
    PrestacaoContas,
    PrestacaoServidor,
    Trecho,
    TrechoRealizado,
    Viajante,
)

DESLOCAMENTO = 1_000_000  # tira as posições do caminho ao reordenar (ordem é única)
PENDENCIA = ("Preencha o km de todos os trechos do diário de bordo, ou anexe o diário "
             "assinado.")


class DiarioInvalido(Exception):
    pass


# ---------------------------------------------------------------- leitura
def equipe_finalizada(prestacao: PrestacaoContas) -> bool:
    linhas = PrestacaoServidor.objects.filter(prestacao=prestacao, removida_em__isnull=True)
    return linhas.exists() and not linhas.filter(finalizada_em__isnull=True).exists()


def pode_editar(usuario, prestacao: PrestacaoContas, finalizada: bool | None = None) -> bool:
    if not policies.pode_editar_equipe_prestacao(usuario, prestacao):
        return False
    return not (equipe_finalizada(prestacao) if finalizada is None else finalizada)


def trechos_do_oficio(oficio: Oficio) -> list[Trecho]:
    return list(oficio.trechos.select_related("origem", "destino").order_by("ordem"))


@transaction.atomic
def obter(prestacao: PrestacaoContas) -> DiarioBordo:
    """O diário da prestação, criado na primeira visita e acertado pelos trechos."""
    diario, _ = DiarioBordo.objects.get_or_create(prestacao=prestacao)
    diario.prestacao = prestacao  # o que a tela já buscou (com o ofício) não se busca de novo
    sincronizar(diario)
    return diario


@transaction.atomic
def sincronizar(diario: DiarioBordo) -> list[DiarioBordoTrecho]:
    """Uma linha por trecho, na ordem; reaproveita a linha do trecho (ou, se o trecho foi
    refeito, a que sobrou na mesma posição) para não perder km digitado."""
    from . import realizado as viagem_realizada
    realizados = viagem_realizada.trechos(diario.prestacao)
    # Com a viagem ajustada (9b-2), as linhas seguem os realizados (referência: o diário
    # segue o roteiro efetivo); a chave de cada linha diz de qual trecho ela é.
    fontes: list[tuple[tuple[str, int], int | None, TrechoRealizado | None]]
    if realizados:
        fontes = [(("r", r.pk), r.trecho_oficio_id, r) for r in realizados]
    else:
        fontes = [(("t", t.pk), t.pk, None) for t in trechos_do_oficio(diario.prestacao.oficio)]

    def chave(lin: DiarioBordoTrecho):
        return ("r", lin.realizado_id) if lin.realizado_id else ("t", lin.trecho_id)

    atuais = list(diario.linhas.order_by("ordem", "pk"))
    if [(chave(lin), lin.ordem) for lin in atuais] == [(f[0], i) for i, f in enumerate(fontes)]:
        return atuais
    diario.linhas.update(ordem=F("ordem") + DESLOCAMENTO)
    existentes = list(diario.linhas.order_by("ordem", "pk"))
    por_chave = {chave(lin): lin for lin in existentes}
    por_trecho = {lin.trecho_id: lin for lin in existentes if lin.trecho_id}
    chaves = {f[0] for f in fontes}
    usados = []
    sobrando = [lin for lin in existentes if chave(lin) not in chaves]
    for i, (k, trecho_id, real) in enumerate(fontes):
        # A mesma chave; senão a linha do mesmo trecho do ofício (ao ajustar ou desfazer o
        # ajuste os km ficam); senão a que sobrou na posição.
        linha = por_chave.get(k)
        if linha is None and trecho_id and por_trecho.get(trecho_id) in sobrando:
            linha = por_trecho[trecho_id]
            sobrando.remove(linha)
        if linha is None and sobrando:
            linha = sobrando.pop(0)
        if linha is None:
            linha = DiarioBordoTrecho(diario=diario, abastecimento=True)
        linha.trecho_id, linha.realizado, linha.ordem = trecho_id, real, i
        linha.save()
        usados.append(linha.pk)
    diario.linhas.exclude(pk__in=usados).delete()
    return list(diario.linhas.order_by("ordem", "pk"))


def linhas(diario: DiarioBordo) -> list[DiarioBordoTrecho]:
    return list(diario.linhas.select_related("trecho__origem", "trecho__destino",
                                             "realizado__origem", "realizado__destino")
                .order_by("ordem", "pk"))


def fonte(linha: DiarioBordoTrecho):
    """De onde vêm rota, horários e distância da linha: o trecho realizado ou o do ofício."""
    return linha.realizado or linha.trecho


def rota(linha: DiarioBordoTrecho) -> str:
    t = fonte(linha)
    return f"{t.origem} → {t.destino}" if t else f"trecho {linha.ordem + 1}"


def prevista(linha: DiarioBordoTrecho) -> int | None:
    t = fonte(linha)
    return round(t.distancia_km) if t and t.distancia_km else None


def _linhas_do_dominio(lista: list[DiarioBordoTrecho]) -> list[dominio.Linha]:
    return [dominio.Linha(rota(lin), lin.km_inicial, lin.km_final, prevista(lin))
            for lin in lista]


def preenchido(prestacao: PrestacaoContas) -> bool:
    """O que a finalização cobra (9b; o diário assinado entra com os anexos, 9d)."""
    diario = DiarioBordo.objects.filter(prestacao=prestacao).first()
    if diario is None:
        return False
    return dominio.preenchido([dominio.Linha("", lin.km_inicial, lin.km_final, None)
                               for lin in diario.linhas.all()])


def preenchidos(prestacao_ids) -> set[int]:
    """As prestações (dentre as dadas) com o diário preenchido, numa consulta só (lista)."""
    from django.db.models import Count
    return set(DiarioBordo.objects.filter(prestacao_id__in=list(prestacao_ids))
               .annotate(n=Count("linhas"),
                         faltam=Count("linhas", filter=Q(linhas__km_inicial__isnull=True)
                                      | Q(linhas__km_final__isnull=True)))
               .filter(n__gt=0, faltam=0).values_list("prestacao_id", flat=True))


def viatura_id(diario: DiarioBordo) -> int | None:
    """A viatura do cadastro usada (a trocada ou a do ofício); a preenchida à mão não conta."""
    if diario.viatura_modo == DiarioBordo.ViaturaModo.CADASTRO:
        return diario.viatura_id
    if diario.viatura_modo == DiarioBordo.ViaturaModo.OFICIO:
        return diario.prestacao.oficio.viatura_id
    return None


def ultimo_km_da_viatura(diario: DiarioBordo) -> dominio.UltimoKm | None:
    """O último km de chegada da mesma viatura em outro diário, de viagem que chegou antes da
    saída desta (só informação: nada é preenchido com ele)."""
    vid = viatura_id(diario)
    if not vid:
        return None
    candidatas = (DiarioBordoTrecho.objects.filter(km_final__isnull=False)
                  .exclude(diario=diario)
                  .filter(Q(diario__viatura_modo=DiarioBordo.ViaturaModo.CADASTRO,
                            diario__viatura_id=vid)
                          | Q(diario__viatura_modo=DiarioBordo.ViaturaModo.OFICIO,
                              diario__prestacao__oficio__viatura_id=vid)))
    # Os horários da linha: os realizados quando a viagem foi ajustada, senão os do ofício.
    candidatas = candidatas.annotate(
        _chegada=Coalesce("realizado__chegada_em", "trecho__chegada_em"))
    saida = (diario.linhas.annotate(_saida=Coalesce("realizado__saida_em", "trecho__saida_em"))
             .filter(_saida__isnull=False).order_by("_saida")
             .values_list("_saida", flat=True).first())
    if saida is not None:
        candidatas = candidatas.filter(Q(_chegada__lte=saida) | Q(_chegada__isnull=True))
    linha = (candidatas.select_related("trecho", "realizado", "diario__prestacao__oficio")
             .order_by(F("_chegada").desc(nulls_last=True), "-km_final", "-pk")
             .first())
    if linha is None:
        return None
    origem = fonte(linha)
    quando = origem.chegada_em if origem else linha.diario.atualizado_em
    return dominio.UltimoKm(linha.km_final or 0, f"{timezone.localtime(quando):%d/%m/%Y}",
                            linha.diario.prestacao.oficio.numero_formatado)


def conferencia(diario: DiarioBordo, lista: list[DiarioBordoTrecho] | None = None
                ) -> dominio.Conferencia:
    lista = linhas(diario) if lista is None else lista
    return dominio.conferir(_linhas_do_dominio(lista), ultimo_km_da_viatura(diario))


# ---------------------------------------------------------------- motorista e viatura
def motorista_do_oficio(oficio: Oficio) -> tuple[str, str]:
    """(nome, CPF) do motorista do ofício: o da equipe ou o de fora."""
    if oficio.motorista_externo == Oficio.MotoristaExterno.SERVIDOR:
        s = oficio.motorista_externo_servidor
        return (s.nome, s.cpf) if s else ("", "")
    if oficio.motorista_externo == Oficio.MotoristaExterno.MANUAL:
        return oficio.motorista_externo_nome.strip(), oficio.motorista_externo_cpf
    v = (Viajante.objects.filter(oficio=oficio, motorista=True).select_related("servidor")
         .first())
    return (v.servidor.nome, v.servidor.cpf) if v else ("", "")


def motorista(diario: DiarioBordo) -> tuple[str, str]:
    if diario.motorista_modo == DiarioBordo.Motorista.SERVIDOR and diario.motorista_servidor:
        return diario.motorista_servidor.nome, diario.motorista_servidor.cpf
    if diario.motorista_modo == DiarioBordo.Motorista.OUTRO:
        return diario.motorista_nome, diario.motorista_cpf
    return motorista_do_oficio(diario.prestacao.oficio)


@dataclass(frozen=True)
class DadosViatura:
    modelo: str
    tipo: str
    placa: str
    combustivel: str

    @property
    def rotulo(self) -> str:
        return f"{self.modelo} ({self.tipo})" if self.modelo and self.tipo else self.modelo


def _da_viatura(v) -> DadosViatura:
    return DadosViatura(v.modelo, v.get_tipo_display() if v.tipo else "", v.placa,
                        str(v.combustivel or ""))


def viatura_do_oficio(oficio: Oficio) -> DadosViatura:
    if oficio.viatura_id:
        return _da_viatura(oficio.viatura)
    return DadosViatura(oficio.transporte_descricao, "", oficio.transporte_placa,
                        str(oficio.transporte_combustivel or ""))


def viatura(diario: DiarioBordo) -> DadosViatura:
    if diario.viatura_modo == DiarioBordo.ViaturaModo.CADASTRO and diario.viatura:
        return _da_viatura(diario.viatura)
    if diario.viatura_modo == DiarioBordo.ViaturaModo.MANUAL:
        tipo = diario.get_viatura_tipo_display() if diario.viatura_tipo else ""
        return DadosViatura(diario.viatura_modelo, str(tipo), diario.viatura_placa,
                            diario.viatura_combustivel)
    return viatura_do_oficio(diario.prestacao.oficio)


def alteracoes(diario: DiarioBordo, *, do_oficio: tuple[str, str] | None = None,
               viatura_oficio: DadosViatura | None = None) -> list[str]:
    """O que o diário trocou em relação ao ofício (vai para o RT, 9c). Quem já calculou o
    motorista e a viatura do ofício passa (a folha), para não consultar de novo."""
    saida = []
    oficio = diario.prestacao.oficio
    if diario.motorista_modo != DiarioBordo.Motorista.OFICIO:
        de = (do_oficio or motorista_do_oficio(oficio))[0]
        para = motorista(diario)[0]
        if de != para:
            saida.append(f"motorista trocado de {de or '—'} para {para or '—'}")
    if diario.viatura_modo != DiarioBordo.ViaturaModo.OFICIO:
        de_v, para_v = viatura_oficio or viatura_do_oficio(oficio), viatura(diario)
        if (de_v.placa, de_v.modelo) != (para_v.placa, para_v.modelo):
            saida.append(f"viatura trocada de {de_v.rotulo or '—'} {de_v.placa}".rstrip()
                         + f" para {para_v.rotulo or '—'} {para_v.placa}".rstrip())
    return saida


# ---------------------------------------------------------------- escrita
def _travar(usuario, diario_pk: int) -> DiarioBordo:
    diario = (DiarioBordo.objects.select_for_update()
              .select_related("prestacao__oficio").get(pk=diario_pk))
    if equipe_finalizada(diario.prestacao):
        raise DiarioInvalido("Prestação finalizada — reabra para editar.")
    policies.exigir(policies.pode_editar_equipe_prestacao(usuario, diario.prestacao),
                    "Você não pode alterar este diário.")
    return diario


@transaction.atomic
def salvar_linhas(usuario, diario_pk: int,
                  valores: dict[int, dict[str, object]]) -> int:
    """{linha_pk: {"km_inicial", "km_final", "abastecimento"}} — uma gravação por linha
    (os dois km se conferem juntos). Devolve quantas mudaram."""
    diario = _travar(usuario, diario_pk)
    por_pk = {lin.pk: lin for lin in diario.linhas.all()}
    mudaram = 0
    for pk, campos in valores.items():
        linha = por_pk.get(pk)
        if linha is None:
            continue
        antes = (linha.km_inicial, linha.km_final, linha.abastecimento)
        if "km_inicial" in campos:
            linha.km_inicial = dominio.numero_do_km(campos["km_inicial"])
        if "km_final" in campos:
            linha.km_final = dominio.numero_do_km(campos["km_final"])
        if "abastecimento" in campos:
            linha.abastecimento = str(campos["abastecimento"] or "") != "nao"
        if (linha.km_inicial, linha.km_final, linha.abastecimento) == antes:
            continue
        try:
            linha.validate_constraints()
        except ValidationError as exc:
            raise DiarioInvalido(f"{rota(linha)}: {' '.join(exc.messages)}") from exc
        linha.save(update_fields=["km_inicial", "km_final", "abastecimento", "atualizado_em"])
        mudaram += 1
    if mudaram:
        diario.save(update_fields=["atualizado_em"])
    return mudaram


OFICIO_DE_REFERENCIA = re.compile(r"^\d{1,5}/\d{4}$")


@transaction.atomic
def trocar_motorista_e_viatura(usuario, diario_pk: int, *, motorista_modo: str,
                               motorista_servidor_id: int | None = None,
                               motorista_nome: str = "", motorista_cpf: str = "",
                               motorista_oficio: str = "", motorista_protocolo: str = "",
                               viatura_modo: str, viatura_id: int | None = None,
                               viatura_modelo: str = "", viatura_placa: str = "",
                               viatura_tipo: str = "", viatura_combustivel: str = ""
                               ) -> DiarioBordo:
    """Só deste diário — o ofício não muda. Validações e limpezas da referência."""
    diario = _travar(usuario, diario_pk)
    erros = []
    if motorista_modo not in DiarioBordo.Motorista.values:
        motorista_modo = DiarioBordo.Motorista.OFICIO
    if viatura_modo not in DiarioBordo.ViaturaModo.values:
        viatura_modo = DiarioBordo.ViaturaModo.OFICIO
    nome = " ".join((motorista_nome or "").split())
    if motorista_modo == DiarioBordo.Motorista.SERVIDOR:
        equipe = set(diario.prestacao.oficio.viajantes.values_list("servidor_id", flat=True))
        if motorista_servidor_id not in equipe:
            erros.append("Selecione um servidor do ofício.")
    elif motorista_modo == DiarioBordo.Motorista.OUTRO:
        if not nome:
            erros.append("Informe o nome do motorista.")
        if motorista_oficio.strip() and not OFICIO_DE_REFERENCIA.match(motorista_oficio.strip()):
            erros.append("Informe o ofício do motorista no formato número/ano (ex.: 15/2026).")
    modelo = " ".join((viatura_modelo or "").split())
    if viatura_modo == DiarioBordo.ViaturaModo.CADASTRO and not viatura_id:
        erros.append("Selecione uma viatura do cadastro.")
    elif viatura_modo == DiarioBordo.ViaturaModo.MANUAL and not modelo:
        erros.append("Informe o modelo da viatura.")
    if erros:
        raise DiarioInvalido(" ".join(erros))
    diario.motorista_modo, diario.viatura_modo = motorista_modo, viatura_modo
    servidor = motorista_modo == DiarioBordo.Motorista.SERVIDOR
    outro = motorista_modo == DiarioBordo.Motorista.OUTRO
    diario.motorista_servidor_id = motorista_servidor_id if servidor else None
    diario.motorista_nome = nome if outro else ""
    diario.motorista_cpf = re.sub(r"\D", "", motorista_cpf or "")[:11] if outro else ""
    diario.motorista_oficio = motorista_oficio.strip()[:16] if outro else ""
    diario.motorista_protocolo = re.sub(r"\D", "", motorista_protocolo or "")[:30] if outro else ""
    cadastro = viatura_modo == DiarioBordo.ViaturaModo.CADASTRO
    manual = viatura_modo == DiarioBordo.ViaturaModo.MANUAL
    diario.viatura_id = viatura_id if cadastro else None
    diario.viatura_modelo = modelo if manual else ""
    diario.viatura_placa = (re.sub(r"[^A-Za-z0-9]", "", viatura_placa or "").upper()[:8]
                            if manual else "")
    from gestao.cadastros.models import Viatura
    diario.viatura_tipo = (viatura_tipo if manual and viatura_tipo in Viatura.Tipo.values
                           else "")
    diario.viatura_combustivel = (" ".join((viatura_combustivel or "").split())[:60]
                                  if manual else "")
    diario.save()
    from . import relatorio  # a troca explica-se no RT (só se o campo estiver vazio)
    relatorio.preencher_informacoes_complementares(diario.prestacao)
    return diario


# ---------------------------------------------------------------- documento
def nome_do_arquivo(diario: DiarioBordo, formato: str) -> str:
    numero = diario.prestacao.oficio.numero_formatado.replace("/", "-")
    return f"DIARIO_DE_BORDO_OFICIO_{numero}.{'pdf' if formato == 'pdf' else 'xlsx'}"


def dados_do_documento(diario: DiarioBordo) -> dict:
    from gestao.cadastros.models import ConfiguracaoInstitucional

    oficio = diario.prestacao.oficio
    config = ConfiguracaoInstitucional.objects.filter(unidade=oficio.unidade).first()
    nome, cpf = motorista(diario)
    v = viatura(diario)
    if diario.motorista_modo == DiarioBordo.Motorista.OUTRO:
        ref_oficio = diario.motorista_oficio or oficio.motorista_oficio_origem
        protocolo = diario.motorista_protocolo or oficio.motorista_protocolo_origem
    else:
        ref_oficio, protocolo = oficio.numero_formatado, oficio.protocolo
    numero_mot, _, ano_mot = (ref_oficio or "").partition("/")
    if len(protocolo or "") == 9:
        protocolo = f"{protocolo[:2]}.{protocolo[2:5]}.{protocolo[5:8]}-{protocolo[8]}"
    itens = []
    for lin in linhas(diario):
        t = fonte(lin)
        saida = timezone.localtime(t.saida_em) if t else None
        chegada = timezone.localtime(t.chegada_em) if t else None
        itens.append({
            "data_saida": f"{saida:%d/%m/%Y}" if saida else "",
            "hora_saida": f"{saida:%H:%M}" if saida else "",
            "km_inicial": dominio.km(lin.km_inicial) if lin.km_inicial is not None else "",
            "data_chegada": f"{chegada:%d/%m/%Y}" if chegada else "",
            "hora_chegada": f"{chegada:%H:%M}" if chegada else "",
            "km_final": dominio.km(lin.km_final) if lin.km_final is not None else "",
            "origem": str(t.origem) if t else "", "destino": str(t.destino) if t else "",
            "abastecimento": ("" if lin.abastecimento is None
                              else "Sim" if lin.abastecimento else "Não"),
        })
    return {
        "titulo": f"Diário de bordo — Ofício {oficio.numero_formatado}",
        "unidade": (config.nome_extenso if config else oficio.unidade.nome).upper(),
        "rodape": config.rodape if config else "",
        "oficio_motorista": numero_mot, "ano": ano_mot, "protocolo": protocolo or "",
        "viatura": v.rotulo.upper(), "placa": v.placa, "combustivel": v.combustivel.upper(),
        "motorista": nome.upper(), "cpf": cpf, "linhas": itens,
    }


def pdf_do_documento(dados: dict) -> bytes:
    import hashlib

    from django.template.loader import render_to_string
    from weasyprint import HTML

    from .documentos.pdf import ASSETS, _recursos, buscar_recurso
    html = render_to_string("viagens/documentos/diario_bordo.html",
                            {"d": dados, **_recursos(False)})
    return HTML(string=html, base_url=str(ASSETS), url_fetcher=buscar_recurso()).write_pdf(
        pdf_variant="pdf/a-2a", pdf_identifier=hashlib.sha256(html.encode()).digest()[:16],
        pdf_tags=True, custom_metadata=True, presentational_hints=True)


def planilha_do_documento(dados: dict) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    from .exportacao import _texto_seguro

    livro = Workbook()
    aba = livro.active or livro.create_sheet()
    aba.title = "Diário de bordo"
    negrito = Font(bold=True)
    aba.append([_texto_seguro(dados["unidade"])])
    aba.append(["DIÁRIO DE BORDO"])
    aba.append([])
    for rotulo, valor in (("Ofício do motorista", f"{dados['oficio_motorista']}/{dados['ano']}"
                           if dados["ano"] else dados["oficio_motorista"]),
                          ("Protocolo", dados["protocolo"]), ("Viatura", dados["viatura"]),
                          ("Placa", dados["placa"]), ("Combustível", dados["combustivel"]),
                          ("Motorista", dados["motorista"]), ("CPF", dados["cpf"])):
        aba.append([rotulo, _texto_seguro(valor)])
        aba.cell(row=aba.max_row, column=1).font = negrito
    aba.append([])
    colunas = ["Data saída", "Hora saída", "Km inicial", "Data chegada", "Hora chegada",
               "Km final", "Origem", "Destino", "Abastecimento"]
    aba.append(colunas)
    for celula in aba[aba.max_row]:
        celula.font = negrito
        celula.alignment = Alignment(horizontal="center")
    chaves = ["data_saida", "hora_saida", "km_inicial", "data_chegada", "hora_chegada",
              "km_final", "origem", "destino", "abastecimento"]
    for item in dados["linhas"]:
        aba.append([_texto_seguro(item[k]) for k in chaves])
    for i, largura in enumerate((12, 9, 11, 12, 9, 11, 26, 26, 14), start=1):
        aba.column_dimensions[get_column_letter(i)].width = largura
    aba.cell(row=1, column=1).font = negrito
    aba.cell(row=2, column=1).font = Font(bold=True, size=13)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
