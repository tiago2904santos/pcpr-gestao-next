"""Diário de bordo (módulo 9b): linhas pelos trechos do ofício (guardam o digitado quando os
trechos mudam), km com validação, motorista e viatura só no diário, conferência do
hodômetro, trava pela equipe finalizada, pendência na finalização, permissões, a folha e o
documento."""

from __future__ import annotations

from io import BytesIO

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from gestao.cadastros.models import Viatura
from gestao.viagens import diario, prestacao
from gestao.viagens.models import DiarioBordo, Oficio, PrestacaoContas, Trecho

from .cenarios import cenario_completo

pytestmark = pytest.mark.django_db


@pytest.fixture
def c():
    return cenario_completo()


def _p(c) -> PrestacaoContas:
    return PrestacaoContas.objects.select_related("oficio").get(
        oficio_id=c.ids["oficio_emitido"])


def _cliente(c, login: str) -> Client:
    cliente = Client()
    cliente.force_login(c.usuarios[login])
    return cliente


def test_uma_linha_por_trecho_e_guarda_o_digitado(c):
    p = _p(c)
    d = diario.obter(p)
    linhas = diario.linhas(d)
    trechos = list(p.oficio.trechos.order_by("ordem"))
    assert [linha.trecho_id for linha in linhas] == [t.pk for t in trechos]
    assert all(linha.abastecimento is True for linha in linhas)  # padrão: Sim
    diario.salvar_linhas(c.usuarios["operador"], d.pk, {
        linhas[0].pk: {"km_inicial": "10.000", "km_final": "10.380"},
        linhas[1].pk: {"km_inicial": "10400", "abastecimento": "nao"}})
    # O trecho é refeito (outra linha no banco): a linha na mesma posição guarda os km.
    primeiro = trechos[0]
    Trecho.objects.filter(pk=primeiro.pk).delete()
    primeiro.pk = None
    primeiro.save()
    diario.sincronizar(d)
    novas = diario.linhas(d)
    assert len(novas) == 2
    por_trecho = {linha.trecho_id: linha for linha in novas}
    assert por_trecho[trechos[1].pk].km_inicial == 10400
    assert por_trecho[trechos[1].pk].abastecimento is False
    assert por_trecho[primeiro.pk].km_final == 10380


def test_km_final_menor_que_o_inicial_nao_grava(c):
    d = diario.obter(_p(c))
    linha = diario.linhas(d)[0]
    with pytest.raises(diario.DiarioInvalido, match="km final não pode ser menor"):
        diario.salvar_linhas(c.usuarios["operador"], d.pk,
                             {linha.pk: {"km_inicial": "500", "km_final": "400"}})
    linha.refresh_from_db()
    assert linha.km_inicial is None


def test_conferencia_e_ultimo_km_da_mesma_viatura(c):
    p = _p(c)
    d = diario.obter(p)
    a, b = diario.linhas(d)
    diario.salvar_linhas(c.usuarios["operador"], d.pk, {
        a.pk: {"km_inicial": 1000, "km_final": 1100}, b.pk: {"km_inicial": 1050,
                                                           "km_final": 1150}})
    conf = diario.conferencia(d)
    assert any("o hodômetro voltou para trás" in aviso for aviso in conf.avisos)
    assert conf.total_rodado == 200


def test_trava_quando_a_equipe_toda_esta_finalizada(c):
    p = _p(c)
    d = diario.obter(p)
    op = c.usuarios["operador"]
    for ps in prestacao.ativos().filter(prestacao=p):
        prestacao.finalizar(op, ps.pk, "Teste da trava.")
    assert diario.equipe_finalizada(p) and not diario.pode_editar(op, p)
    with pytest.raises(diario.DiarioInvalido, match="reabra para editar"):
        diario.salvar_linhas(op, d.pk, {})


def test_pendencia_do_diario_na_finalizacao(c):
    p = _p(c)
    ps = prestacao.ativos().filter(prestacao=p).first()
    assert diario.PENDENCIA in prestacao.pendencias(ps)
    d = diario.obter(p)
    diario.salvar_linhas(c.usuarios["operador"], d.pk, {
        linha.pk: {"km_inicial": 1000 + i * 400, "km_final": 1300 + i * 400}
        for i, linha in enumerate(diario.linhas(d))})
    assert diario.PENDENCIA not in prestacao.pendencias(ps)
    assert diario.preenchidos([p.pk]) == {p.pk}


def test_motorista_e_viatura_so_neste_diario(c):
    p = _p(c)
    d = diario.obter(p)
    op = c.usuarios["operador"]
    oficio = p.oficio
    antes = (oficio.viatura_id, list(oficio.viajantes.filter(motorista=True)
                                     .values_list("servidor_id", flat=True)))
    with pytest.raises(diario.DiarioInvalido, match="servidor que dirigiu"):
        diario.trocar_motorista_e_viatura(op, d.pk, motorista_modo="servidor",
                                          motorista_servidor_id=None, viatura_modo="oficio")
    with pytest.raises(diario.DiarioInvalido, match="ofício pelo qual o motorista"):
        diario.trocar_motorista_e_viatura(op, d.pk, motorista_modo="outro_oficio",
                                          viatura_modo="oficio")
    with pytest.raises(diario.DiarioInvalido, match="modelo da viatura"):
        diario.trocar_motorista_e_viatura(op, d.pk, motorista_modo="oficio",
                                          viatura_modo="manual")
    outra = Viatura.objects.exclude(pk=oficio.viatura_id).first()
    colega = oficio.viajantes.filter(motorista=False).select_related("servidor").first().servidor
    diario.trocar_motorista_e_viatura(
        op, d.pk, motorista_modo="servidor", motorista_servidor_id=colega.pk,
        viatura_modo="cadastro", viatura_id=outra.pk)
    d.refresh_from_db()
    assert diario.motorista(d) == (colega.nome, colega.cpf)
    assert diario.viatura(d).placa == outra.placa
    alteracoes = diario.alteracoes(d)
    assert len(alteracoes) == 2 and colega.nome in alteracoes[0]
    oficio.refresh_from_db()
    assert (oficio.viatura_id, list(oficio.viajantes.filter(motorista=True).values_list(
        "servidor_id", flat=True))) == antes  # o ofício não muda
    # Voltar para o do ofício limpa o que não vale mais.
    diario.trocar_motorista_e_viatura(op, d.pk, motorista_modo="oficio", viatura_modo="oficio")
    d.refresh_from_db()
    assert d.motorista_servidor_id is None and d.viatura_id is None
    assert diario.alteracoes(d) == []


def test_motorista_de_outro_oficio_usa_o_diario_de_la(c):
    """Escolhido o ofício pelo qual o motorista viajou, o diário que vale é o da prestação
    daquele ofício: motorista, viatura, km, pendência e documento vêm de lá; o RT conta a
    troca com motorista, viatura, ofício e protocolo."""
    p = _p(c)
    op = c.usuarios["operador"]
    d = diario.obter(p)
    de_la = Oficio.objects.get(pk=c.ids["oficio_rascunho"])

    def trocar(oficio_id, **extra):
        diario.trocar_motorista_e_viatura(op, d.pk, motorista_modo="outro_oficio",
                                          oficio_do_motorista_id=oficio_id,
                                          viatura_modo="oficio", **extra)

    with pytest.raises(diario.DiarioInvalido, match="ainda não tem prestação"):
        trocar(de_la.pk)
    with pytest.raises(diario.DiarioInvalido, match="própria prestação"):
        trocar(p.oficio_id)
    # A prestação só existe para ofício emitido (prestacao.sincronizar): o de lá também é.
    Oficio.objects.filter(pk=de_la.pk).update(situacao=Oficio.Situacao.EMITIDO,
                                              emitido_em=timezone.now())
    p2 = PrestacaoContas.objects.create(oficio=de_la)
    d2 = diario.obter(p2)
    diario.salvar_linhas(op, d2.pk, {linha.pk: {"km_inicial": 500 + i * 100,
                                                "km_final": 580 + i * 100}
                                     for i, linha in enumerate(diario.linhas(d2))})
    trocar(de_la.pk, viatura_id=Viatura.objects.first().pk)
    d.refresh_from_db()
    assert d.oficio_do_motorista_id == de_la.pk and d.viatura_modo == "oficio"
    assert diario.vigente(d).pk == d2.pk
    assert diario.motorista(d) == diario.motorista(d2)
    assert diario.viatura(d) == diario.viatura(d2)
    # Os km são os de lá: este diário fica preenchido sem digitar nada aqui.
    assert diario.preenchido(p) and diario.preenchidos([p.pk]) == {p.pk}
    assert diario.dados_do_documento(d)["linhas"][0]["km_inicial"] == "500"
    nome, _ = diario.motorista(d2)
    [texto] = diario.alteracoes(d)
    assert texto.startswith("houve troca de motorista") and nome in texto
    assert f"Ofício {de_la.numero_formatado}" in texto
    # O diário de lá não pode, por sua vez, apontar para um terceiro.
    with pytest.raises(diario.DiarioInvalido, match="também usa o de outro"):
        diario.trocar_motorista_e_viatura(op, d2.pk, motorista_modo="outro_oficio",
                                          oficio_do_motorista_id=p.oficio_id,
                                          viatura_modo="oficio")


def test_permissoes(c):
    p = _p(c)
    outra, consulta = _cliente(c, "outra"), _cliente(c, "consulta")
    assert outra.get(reverse("viagens:diario", args=[p.pk])).status_code == 404
    assert outra.post(reverse("viagens:autosave_diario", args=[p.pk])).status_code == 404
    r = consulta.get(reverse("viagens:diario", args=[p.pk]))
    assert r.status_code == 200 and 'name="l-' not in r.content.decode()
    assert not DiarioBordo.objects.filter(prestacao=p).exists()  # consulta não cria
    r = consulta.post(reverse("viagens:autosave_diario", args=[p.pk]))
    assert r.json()["salvo"] is False


def test_folha_autosave_e_documento(c, django_assert_max_num_queries):
    p = _p(c)
    op = _cliente(c, "operador")
    r = op.get(reverse("viagens:diario", args=[p.pk]))
    html = r.content.decode()
    assert r.status_code == 200 and 'id="conferencia"' in html
    assert "Preencher pelas distâncias" in html and "Ir para Relatório técnico" in html
    d = DiarioBordo.objects.get(prestacao=p)
    a, b = diario.linhas(d)
    assert f'name="l-{a.pk}-km_inicial"' in html
    r = op.post(reverse("viagens:autosave_diario", args=[p.pk]), {
        f"l-{a.pk}-km_inicial": "20.000", f"l-{a.pk}-km_final": "20.400",
        f"l-{b.pk}-km_inicial": "20.410", f"l-{b.pk}-km_final": "20.300"})
    assert r.json() == {"salvo": False, "mensagem": f"{diario.rota(b)}: O km final não pode "
                                                    "ser menor que o km inicial."}
    r = op.post(reverse("viagens:autosave_diario", args=[p.pk]), {
        f"l-{a.pk}-km_inicial": "20.000", f"l-{a.pk}-km_final": "20.400",
        f"l-{b.pk}-abastecimento": "nao"})
    assert r.json()["salvo"] is True
    a.refresh_from_db()
    assert (a.km_inicial, a.km_final) == (20000, 20400)
    with django_assert_max_num_queries(45):
        assert op.get(reverse("viagens:diario", args=[p.pk])).status_code == 200
    r = op.get(reverse("viagens:documento_diario", args=[p.pk, "xlsx"]))
    assert "DIARIO_DE_BORDO_OFICIO_" in r["Content-Disposition"]
    folha = load_workbook(BytesIO(r.content)).active
    valores = [str(v) for linha in folha.iter_rows(values_only=True) for v in linha if v]
    assert "20.000" in valores and "DIÁRIO DE BORDO" in valores
    r = op.get(reverse("viagens:documento_diario", args=[p.pk, "pdf"]))
    assert r["Content-Type"] == "application/pdf" and r.content[:4] == b"%PDF"
    assert op.get(reverse("viagens:documento_diario", args=[p.pk, "doc"])).status_code == 404
    # A lista da prestação leva ao diário.
    assert reverse("viagens:diario", args=[p.pk]) in op.get(
        reverse("viagens:prestacoes")).content.decode()


def test_troca_pela_tela(c):
    p = _p(c)
    op = _cliente(c, "operador")
    r = op.post(reverse("viagens:motorista_diario", args=[p.pk]),
                {"motorista_modo": "outro_oficio", "oficio_do_motorista": "",
                 "viatura_modo": "oficio"}, follow=True)
    assert "Busque e escolha o ofício pelo qual o motorista viajou." in r.content.decode()
    r = op.post(reverse("viagens:motorista_diario", args=[p.pk]),
                {"motorista_modo": "oficio", "viatura_modo": "manual",
                 "viatura_modelo": "Hilux", "viatura_placa": "abc-1d23",
                 "viatura_tipo": "caracterizada"})
    assert r["Location"].endswith("#motorista")
    d = DiarioBordo.objects.get(prestacao=p)
    assert (d.viatura_modelo, d.viatura_placa) == ("Hilux", "ABC1D23")
    assert Oficio.objects.get(pk=p.oficio_id).viatura_id is not None


def test_motorista_e_transporte_gravam_sozinhos(c):
    p = _p(c)
    op = _cliente(c, "operador")
    url = reverse("viagens:autosave_motorista_diario", args=[p.pk])
    # Escolha sem o que ela pede: avisa e não grava.
    r = op.post(url, {"motorista_modo": "servidor", "viatura_modo": "oficio"}).json()
    assert r == {"salvo": False, "mensagem": "Selecione o servidor que dirigiu."}
    r = op.post(url, {"motorista_modo": "oficio", "viatura_modo": "manual",
                      "viatura_modelo": "Hilux"}).json()
    assert r["salvo"] is True
    d = DiarioBordo.objects.get(prestacao=p)
    assert (d.viatura_modo, d.viatura_modelo) == ("manual", "Hilux")
    html = op.get(reverse("viagens:diario", args=[p.pk])).content.decode()
    assert 'id="viatura-atual"' in html and "informada à mão, só neste diário" in html
    assert "Aplicar motorista e viatura" not in html
    consulta = _cliente(c, "consulta")
    assert consulta.post(url, {"motorista_modo": "oficio", "viatura_modo": "oficio"}).json()[
        "salvo"] is False
