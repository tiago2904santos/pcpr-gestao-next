"""Prestação de contas, base (módulo 9a): nascimento pela emissão do ofício, equipe que muda,
lote de solicitações, arquivar/finalizar (servidor e equipe), envio/aprovação/devolução,
abas, permissões por unidade e a tela."""

from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO

import pytest
from django.test import Client
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from gestao.plataforma.models import Notificacao
from gestao.viagens import prestacao
from gestao.viagens.models import Oficio, PrestacaoContas, PrestacaoServidor, Viajante

from .cenarios import cenario_completo

pytestmark = pytest.mark.django_db


@pytest.fixture
def c():
    return cenario_completo()


def _linhas(c) -> list[PrestacaoServidor]:
    return list(PrestacaoServidor.objects.filter(
        prestacao__oficio_id=c.ids["oficio_emitido"]).select_related("servidor")
        .order_by("servidor__nome"))


def _cliente(c, login: str) -> Client:
    cliente = Client()
    cliente.force_login(c.usuarios[login])
    return cliente


def _preencher_diario(c, ps):
    """O diário da equipe com km em todos os trechos (a finalização cobra)."""
    from gestao.viagens import diario
    d = diario.obter(ps.prestacao)
    km = 10000
    valores = {}
    for linha in diario.linhas(d):
        valores[linha.pk] = {"km_inicial": km, "km_final": km + 300}
        km += 300
    diario.salvar_linhas(c.usuarios["operador"], d.pk, valores)


def _preencher_relatorio(c, ps):
    """O relatório técnico da equipe com descrição, objetivo e conclusão."""
    from gestao.viagens import relatorio
    rt = relatorio.obter(ps.prestacao)
    relatorio.salvar(c.usuarios["operador"], rt.pk, {
        "motivo": "Evento (teste).", "atividade": "Apoio (teste).",
        "conclusao": "Concluído (teste)."})


def _preencher(c, ps, numero="2026/0001", liberacao=date(2030, 1, 6),
               prazo=date(2030, 1, 9)):
    _preencher_diario(c, ps)
    _preencher_relatorio(c, ps)
    return prestacao.salvar_solicitacao(c.usuarios["operador"], ps.pk, numero=numero,
                                        liberacao=liberacao, prazo=prazo)


# ------------------------------------------------------------------ nascimento
def test_emitir_cria_a_prestacao_com_a_equipe(c):
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    p = PrestacaoContas.objects.get(oficio=oficio)
    equipe = set(oficio.viajantes.values_list("servidor_id", flat=True))
    assert set(p.servidores.values_list("servidor_id", flat=True)) == equipe
    assert len(equipe) == 2
    # Rascunho e cancelado não têm prestação.
    assert not PrestacaoContas.objects.filter(oficio_id__in=[
        c.ids["oficio_rascunho"], c.ids["oficio_cancelado"]]).exists()
    # Sincronizar de novo não duplica.
    prestacao.sincronizar(oficio)
    assert p.servidores.count() == 2


def test_quem_sai_da_equipe_some_ou_fica_guardado_com_dados(c):
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    a, b = _linhas(c)
    _preencher(c, a)
    Viajante.objects.filter(oficio=oficio).delete()
    prestacao.sincronizar(oficio)
    assert not PrestacaoServidor.objects.filter(pk=b.pk).exists()  # sem dados: apagada
    a.refresh_from_db()
    assert a.removida_em is not None  # com dados: guardada, fora das listas
    assert not prestacao.ativos().filter(pk=a.pk).exists()
    # Voltou para a equipe: a linha guardada volta com os dados.
    Viajante.objects.create(oficio=oficio, servidor=a.servidor, ordem=1)
    prestacao.sincronizar(oficio)
    a.refresh_from_db()
    assert a.removida_em is None and a.numero_solicitacao == "2026/0001"


# ------------------------------------------------------------------ solicitação e trava
def _colega_da_unidade(c):
    from django.contrib.auth.models import Group

    from gestao.cadastros.models import Lotacao
    from gestao.identidade.models import Usuario
    u = Usuario.objects.create_user("colega", "colega@pc.pr.gov.br", "senha-local-123",
                                    nome="Colega da ASCOM")
    u.groups.add(Group.objects.get(name="OPERADOR_VIAGENS"))
    Lotacao.objects.create(usuario=u, unidade=c.usuarios["operador"].lotacao.unidade)
    return u


def test_solicitacao_valida_datas_e_avisa_a_unidade(c):
    _colega_da_unidade(c)
    a, _ = _linhas(c)
    with pytest.raises(prestacao.PrestacaoInvalida, match="anterior à liberação"):
        _preencher(c, a, liberacao=date(2030, 1, 9), prazo=date(2030, 1, 6))
    _preencher(c, a)
    a.refresh_from_db()
    assert a.situacao == PrestacaoServidor.Situacao.PREENCHIMENTO
    avisados = set(Notificacao.objects.filter(titulo__startswith="Diárias liberadas")
                   .values_list("usuario__login", flat=True))
    # Só a equipe de viagens da unidade do ofício, menos quem lançou.
    assert avisados == {"colega"}
    # Liberar de novo (já liberada) não repete o aviso.
    n = Notificacao.objects.count()
    _preencher(c, a, numero="2026/0002")
    assert Notificacao.objects.count() == n


def test_finalizar_pede_justificativa_com_pendencia_e_trava(c):
    a, b = _linhas(c)
    op = c.usuarios["operador"]
    with pytest.raises(prestacao.PrestacaoInvalida, match="justificativa"):
        prestacao.finalizar(op, a.pk)
    prestacao.finalizar(op, a.pk, "Solicitação ainda no papel; segue depois.")
    a.refresh_from_db()
    assert a.finalizada and a.justificativa_finalizacao
    with pytest.raises(prestacao.PrestacaoInvalida, match="reabra para editar"):
        _preencher(c, a)
    prestacao.reabrir(op, a.pk)
    a.refresh_from_db()
    assert not a.finalizada and a.justificativa_finalizacao == ""
    # Sem pendência finaliza direto e não guarda justificativa.
    _preencher(c, b)
    prestacao.finalizar(op, b.pk, "não precisava")
    b.refresh_from_db()
    assert b.finalizada and b.justificativa_finalizacao == ""


def test_equipe_finaliza_quem_nao_tem_pendencia(c):
    a, b = _linhas(c)
    _preencher(c, a)
    r = prestacao.acao_da_equipe(c.usuarios["operador"], a.prestacao_id, "finalizar")
    assert r.feitos == 1 and r.pulados == [b.servidor.nome]
    r = prestacao.acao_da_equipe(c.usuarios["operador"], a.prestacao_id, "arquivar")
    assert r.feitos == 2


def test_envio_aprovacao_e_devolucao(c):
    a, _ = _linhas(c)
    op, gestor = c.usuarios["operador"], c.usuarios["gestor"]
    with pytest.raises(prestacao.PrestacaoInvalida, match="Finalize"):
        prestacao.enviar(op, a.pk, enviada_em=None, protocolo="")
    _preencher(c, a)
    prestacao.finalizar(op, a.pk)
    with pytest.raises(prestacao.PrestacaoInvalida, match="enviada"):
        prestacao.aprovar(gestor, a.pk)
    assert prestacao.enviar(op, a.pk, enviada_em=date(2030, 1, 10), protocolo="E-1") == 1
    a.refresh_from_db()
    assert a.situacao == "enviada" and a.protocolo_envio == "E-1"
    with pytest.raises(prestacao.PrestacaoInvalida, match="motivo"):
        prestacao.devolver(gestor, a.pk, "  ")
    # Enviada não reabre por baixo: corrige-se devolvendo.
    with pytest.raises(prestacao.PrestacaoInvalida, match="devolva"):
        prestacao.reabrir(op, a.pk)
    prestacao.aprovar(gestor, a.pk)
    # Aprovada não volta a "enviada" nem reabre (nem pela equipe).
    with pytest.raises(prestacao.PrestacaoInvalida, match="aprovada"):
        prestacao.enviar(op, a.pk, enviada_em=None, protocolo="E-2")
    with pytest.raises(prestacao.PrestacaoInvalida, match="por enviar"):
        prestacao.enviar(op, a.pk, enviada_em=None, protocolo="E-2", equipe=True)
    assert prestacao.acao_da_equipe(op, a.prestacao_id, "reabrir").feitos == 0
    prestacao.devolver(gestor, a.pk, "Comprovante ilegível.")
    a.refresh_from_db()
    assert a.situacao == "devolvida" and not a.finalizada
    assert Notificacao.objects.filter(usuario=op, titulo__startswith="Prestação devolvida",
                                      link__contains="?q=",
                                      mensagem="Comprovante ilegível.").exists()
    assert prestacao.filtrar(PrestacaoServidor.objects.all(), "devolvidas").filter(
        pk=a.pk).exists()


def test_abas(c):
    a, b = _linhas(c)
    todas = prestacao.ativos()
    assert set(prestacao.filtrar(todas, "nao_liberadas")) == {a, b}
    assert set(prestacao.filtrar(todas, "sem_solicitacao")) == {a, b}
    hoje = timezone.localdate()
    _preencher(c, a, liberacao=hoje - timedelta(days=20), prazo=hoje - timedelta(days=10))
    assert set(prestacao.filtrar(todas, "liberadas")) == {a}
    assert set(prestacao.filtrar(todas, "prestacao_vencida")) == {a}
    assert set(prestacao.filtrar(todas, "saque_vencendo")) == {a}
    prestacao.finalizar(c.usuarios["operador"], a.pk)
    # Finalizados: só quando ninguém da prestação está em aberto (arquivar não conta).
    assert not prestacao.filtrar(todas, "finalizados").exists()
    prestacao.arquivar(c.usuarios["operador"], b.pk)
    assert not prestacao.filtrar(todas, "finalizados").exists()
    assert set(prestacao.filtrar(todas, "arquivados")) == {b}
    prestacao.finalizar(c.usuarios["operador"], b.pk, "Não viajou.")
    assert set(prestacao.filtrar(todas, "finalizados")) == {a, b}
    assert set(prestacao.filtrar(todas, "finalizadas_mes")) == {a, b}


def test_diaria_liberada_e_a_por_servidor_do_calculo(c):
    from decimal import Decimal
    a, _ = _linhas(c)
    oficio = a.prestacao.oficio
    assert oficio.diarias_total > 0
    assert prestacao.diaria_liberada(a) == Decimal(oficio.diarias_calculo["por_servidor"])
    # O override (recebida) não muda o teto liberado.
    a.diaria_valor_override = Decimal("1.00")
    assert prestacao.diaria_liberada(a) == Decimal(oficio.diarias_calculo["por_servidor"])
    # Sem o cálculo guardado: o total ÷ a equipe.
    oficio.diarias_calculo = {}
    assert prestacao.diaria_liberada(a) == (oficio.diarias_total / 2).quantize(Decimal("0.01"))


# ------------------------------------------------------------------ permissões
def test_cancelar_o_oficio_tira_da_lista_e_trava(c):
    from gestao.viagens import policies, services
    a, _ = _linhas(c)
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    services.cancelar(oficio, c.usuarios["gestor"], "Evento cancelado (teste).")
    assert not policies.prestacoes_visiveis(c.usuarios["operador"]).exists()
    a.refresh_from_db()
    assert not policies.pode_editar_prestacao(c.usuarios["operador"], a)


def test_linha_removida_e_oficio_reaberto_nao_se_alteram(c):
    from gestao.viagens import policies, services
    a, b = _linhas(c)
    op = c.usuarios["operador"]
    _preencher(c, a)
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    Viajante.objects.filter(oficio=oficio, servidor=a.servidor).delete()
    prestacao.sincronizar(oficio)
    a.refresh_from_db()
    assert a.removida_em is not None
    _cliente(c, "operador").post(reverse("viagens:acao_prestacao", args=[a.pk, "arquivar"]))
    a.refresh_from_db()
    assert not a.arquivada  # quem saiu da equipe não se altera pelo pk antigo
    services.reabrir(oficio, c.usuarios["gestor"], "Corrigir a equipe (teste).")
    b.refresh_from_db()
    assert not policies.prestacoes_visiveis(op).exists()
    assert not policies.pode_editar_prestacao(op, b)


def test_outra_unidade_e_consulta(c):
    a, _ = _linhas(c)
    outra, consulta = _cliente(c, "outra"), _cliente(c, "consulta")
    lista = outra.get(reverse("viagens:prestacoes"))
    assert lista.status_code == 200 and a.servidor.nome not in lista.content.decode()
    r = outra.post(reverse("viagens:acao_prestacao", args=[a.pk, "arquivar"]))
    assert r.status_code == 404
    for rota in ("viagens:salvar_prestacao", "viagens:autosave_prestacao"):
        assert outra.post(reverse(rota, args=[a.pk]), {"numero": "X"}).status_code == 404
    r = outra.post(reverse("viagens:acao_equipe_prestacao", args=[a.prestacao_id, "finalizar"]))
    assert r.status_code == 404
    planilha = load_workbook(BytesIO(outra.get(reverse("viagens:exportar_prestacoes")).content))
    assert a.servidor.nome not in str(list(planilha.active.iter_rows(values_only=True)))
    # Consulta vê sem campos nem ações e não altera nada (nem pela equipe).
    html = consulta.get(reverse("viagens:prestacoes")).content.decode()
    assert a.servidor.nome in html and f'id="ps-{a.pk}-numero"' not in html
    assert "Finalizar" not in html
    r = consulta.post(reverse("viagens:acao_equipe_prestacao", args=[a.prestacao_id,
                                                                     "arquivar"]), follow=True)
    assert "não pode alterar" in r.content.decode()
    assert consulta.post(reverse("viagens:autosave_prestacao", args=[a.pk]),
                         {"numero": "X"}).json()["salvo"] is False
    r = consulta.post(reverse("viagens:acao_prestacao", args=[a.pk, "arquivar"]),
                      {"voltar": "/viagens/prestacoes/"})
    a.refresh_from_db()
    assert r.status_code == 302 and not a.arquivada


def test_anonimo_vai_para_a_entrada(client):
    r = client.get(reverse("viagens:prestacoes"))
    assert r.status_code == 302 and "/entrar" in r["Location"]


# ------------------------------------------------------------------ tela
def test_lista_cartao_e_acoes(c, django_assert_max_num_queries):
    a, b = _linhas(c)
    _preencher_diario(c, a)
    _preencher_relatorio(c, a)
    op = _cliente(c, "operador")
    with django_assert_max_num_queries(30):
        r = op.get(reverse("viagens:prestacoes"))
    html = r.content.decode()
    assert r.status_code == 200
    assert a.servidor.nome in html and b.servidor.nome in html
    assert f'id="ps-{a.pk}-numero"' in html
    assert reverse("viagens:autosave_prestacao", args=[a.pk]) in html
    # Autosave do cartão (contrato do autosave.js).
    r = op.post(reverse("viagens:autosave_prestacao", args=[a.pk]),
                {"numero": "2026/0099", "liberacao": "06/01/2030", "prazo": "09/01/2030"})
    assert r.json()["salvo"] is True
    a.refresh_from_db()
    assert a.numero_solicitacao == "2026/0099" and a.prazo_limite_saque == date(2030, 1, 9)
    r = op.post(reverse("viagens:autosave_prestacao", args=[b.pk]),
                {"numero": "1", "liberacao": "31/02/2030", "prazo": ""})
    assert r.json() == {"salvo": False,
                        "mensagem": "Data de liberação inválida: 31/02/2030. Use dd/mm/aaaa."}
    # Sem JavaScript (Enter): grava e volta para o cartão.
    r = op.post(reverse("viagens:salvar_prestacao", args=[b.pk]),
                {"numero": "2026/0100", "liberacao": "", "prazo": "",
                 "voltar": "/viagens/prestacoes/?aba=liberadas"})
    assert r["Location"] == f"/viagens/prestacoes/?aba=liberadas#ps-{b.pk}"
    b.refresh_from_db()
    assert b.numero_solicitacao == "2026/0100"
    # A ação leva o que está digitado: o prazo entra e a pendência some antes de finalizar.
    op.post(reverse("viagens:acao_prestacao", args=[b.pk, "finalizar"]),
            {"numero": "2026/0100", "liberacao": "", "prazo": "09/01/2030"})
    b.refresh_from_db()
    assert b.finalizada and b.prazo_limite_saque == date(2030, 1, 9)
    assert b.justificativa_finalizacao == ""
    prestacao.reabrir(c.usuarios["operador"], b.pk)
    PrestacaoServidor.objects.filter(pk=b.pk).update(prazo_limite_saque=None)
    b.refresh_from_db()
    # Finalizar pela tela; a aba Liberadas mostra só quem tem liberação.
    op.post(reverse("viagens:acao_prestacao", args=[a.pk, "finalizar"]))
    a.refresh_from_db()
    assert a.finalizada
    r = op.get(reverse("viagens:prestacoes") + "?aba=finalizadas_mes")
    assert [linha["ps"].pk for bloco in r.context["blocos"] for linha in bloco["linhas"]] == [a.pk]
    # Com pendência, sem justificativa, não finaliza.
    r = op.post(reverse("viagens:acao_prestacao", args=[b.pk, "finalizar"]), follow=True)
    b.refresh_from_db()
    assert not b.finalizada and "pendências" in r.content.decode()
    r = op.post(reverse("viagens:acao_equipe_prestacao", args=[a.prestacao_id, "finalizar"]),
                follow=True)
    texto = r.content.decode()
    assert "Ninguém foi finalizado" in texto and b.servidor.nome in texto
    # Envio, aprovação pela gestora.
    op.post(reverse("viagens:acao_prestacao", args=[a.pk, "enviar"]),
            {"enviada_em": "10/01/2030", "protocolo": "E-77"})
    a.refresh_from_db()
    assert a.situacao == "enviada" and a.enviada_em == date(2030, 1, 10)
    _cliente(c, "gestor").post(reverse("viagens:acao_prestacao", args=[a.pk, "aprovar"]))
    a.refresh_from_db()
    assert a.situacao == "aprovada"
    assert op.post(reverse("viagens:acao_prestacao", args=[a.pk, "xyz"])).status_code == 404


def test_busca_por_numero_do_oficio(c):
    a, _ = _linhas(c)
    oficio = a.prestacao.oficio
    op = _cliente(c, "operador")
    html = op.get(reverse("viagens:prestacoes"), {"q": oficio.numero_formatado}).content.decode()
    assert a.servidor.nome in html
    html = op.get(reverse("viagens:prestacoes"), {"q": "ninguém-com-este-nome"}).content.decode()
    assert "Nenhuma prestação encontrada" in html


def test_voltar_externo_e_ignorado(c):
    a, _ = _linhas(c)
    r = _cliente(c, "operador").post(reverse("viagens:acao_prestacao", args=[a.pk, "arquivar"]),
                                     {"voltar": "https://exemplo.invalido/"})
    assert r["Location"] == reverse("viagens:prestacoes") + f"#ps-{a.pk}"


def test_exportar_planilha(c):
    a, _ = _linhas(c)
    r = _cliente(c, "operador").get(reverse("viagens:exportar_prestacoes"))
    assert r.status_code == 200 and "spreadsheetml" in r["Content-Type"]
    folha = load_workbook(BytesIO(r.content)).active
    valores = [str(v) for linha in folha.iter_rows(values_only=True) for v in linha if v]
    assert a.servidor.nome in valores


# ------------------------------------------------------------------ abas dos outros módulos
def _finalizar_todos(c):
    for ps in _linhas(c):
        prestacao.finalizar(c.usuarios["operador"], ps.pk, "Teste das abas.")


def test_contas_prestadas_nos_outros_modulos(c):
    from gestao.viagens import ordens, queries, termos, viagem
    from gestao.viagens.models import OrdemServico, Roteiro, TermoAutorizacao, Viagem

    op = c.usuarios["operador"]
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    Oficio.objects.filter(pk=oficio.pk).update(roteiro_id=c.ids["roteiro"])
    v = viagem.criar(op)
    Oficio.objects.filter(pk=oficio.pk).update(viagem=v)
    ordem, _ = ordens.salvar(op, oficios=[oficio])
    from gestao.cadastros.models import Municipio
    londrina = [Municipio.objects.get(nome="Londrina", uf="PR")]
    termo = termos.salvar(op, oficio=oficio, destinos=londrina)
    avulso = termos.salvar(op, evento="Avulso (teste)", destinos=londrina,
                           data_inicio=date(2030, 5, 10))

    def finalizados():
        return {
            "roteiro": set(queries.filtrar_roteiros(Roteiro.objects.all(), "finalizados")
                           .values_list("pk", flat=True)),
            "oficio": set(Oficio.objects.filter(queries.FILTROS_SITUACAO["prestadas"][1])
                          .values_list("pk", flat=True)),
            "ordem": set(OrdemServico.objects.filter(prestacao.prestadas("ordem"))
                         .values_list("pk", flat=True)),
            "termo": set(TermoAutorizacao.objects.filter(prestacao.prestadas("termo"))
                         .values_list("pk", flat=True)),
            "viagem": set(Viagem.objects.filter(prestacao.prestadas("viagem"))
                          .values_list("pk", flat=True)),
        }

    assert finalizados() == {k: set() for k in ("roteiro", "oficio", "ordem", "termo", "viagem")}
    a, b = _linhas(c)
    prestacao.finalizar(op, a.pk, "Só um.")
    prestacao.arquivar(op, b.pk)  # arquivar não conta como finalizar
    assert not any(finalizados().values())
    prestacao.finalizar(op, b.pk, "O outro.")
    assert finalizados() == {"roteiro": {c.ids["roteiro"]}, "oficio": {oficio.pk},
                             "ordem": {ordem.pk}, "termo": {termo.pk}, "viagem": {v.pk}}
    assert avulso.pk not in finalizados()["termo"]
    # Reabrir uma devolve o registro às abas de quando.
    prestacao.reabrir(op, a.pk)
    assert not any(finalizados().values())


def test_abas_de_quando_excluem_as_contas_prestadas(c):
    from gestao.viagens import viagem
    from gestao.viagens.models import Viagem

    op = c.usuarios["operador"]
    v = viagem.criar(op)
    Viagem.objects.filter(pk=v.pk).update(data_inicio=timezone.localdate() - timedelta(days=3))
    Oficio.objects.filter(pk=c.ids["oficio_emitido"]).update(viagem=v)
    cliente = _cliente(c, "operador")

    def na_aba(aba):
        return v.pk in [linha["v"].pk for linha in cliente.get(
            reverse("viagens:viagens"), {"aba": aba}).context["linhas"]]

    assert na_aba("atuais") and not na_aba("prestadas")
    _finalizar_todos(c)
    assert na_aba("prestadas") and not na_aba("atuais")
    html = cliente.get(reverse("viagens:viagens")).content.decode()
    assert "Contas prestadas" in html
    for rota, rotulo in (("viagens:roteiros", "Finalizados"), ("viagens:termos", "Finalizados"),
                         ("viagens:ordens", "Finalizadas"), ("viagens:planos", "Finalizados"),
                         ("viagens:oficios", "Contas prestadas")):
        assert rotulo in cliente.get(reverse(rota)).content.decode(), rota
    oficios = cliente.get(reverse("viagens:oficios"), {"situacao": "prestadas"})
    assert Oficio.objects.get(pk=c.ids["oficio_emitido"]).numero_formatado in (
        oficios.content.decode())
