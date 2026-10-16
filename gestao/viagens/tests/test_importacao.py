"""Corpus sintético: leitura real, vínculos, isolamento e repetição da fila."""

import base64
import io
from datetime import timedelta
from html import escape

import pytest
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from gestao.plataforma import outbox
from gestao.viagens import assinados, importacao, termos
from gestao.viagens.documentos.importador import LeituraInvalida, analisar
from gestao.viagens.models import AnexoPrestacao, ArquivoImportado, Oficio, ViaAssinada
from gestao.viagens.tests.cenarios import cenario_completo


def pdf_textos(textos):
    from weasyprint import HTML

    partes = [f"<section><pre>{escape(t)}</pre></section>" for t in textos]
    return HTML(
        string='<meta charset="utf-8"><style>section{break-after:page}'
        "pre{white-space:pre-wrap;font:12pt sans-serif}</style>" + "".join(partes)
    ).write_pdf()


@pytest.fixture
def c(db, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    dados = cenario_completo()
    while outbox.processar_lote():
        pass
    return dados


def test_conversao_texto_real():
    texto = "DESPACHO\nOfício 120/2026\nEncaminho para providências administrativas cabíveis."
    r, pdf, partes = analisar(texto.encode(), "arquivo.txt")
    assert pdf.startswith(b"%PDF") and len(partes) == 1
    assert r["partes"][0]["tipo"] == "despacho" and r["ocr"] == 0


def test_conversao_docx_real():
    from docx import Document

    doc = Document()
    doc.add_paragraph("DESPACHO")
    doc.add_paragraph("Referente ao Ofício 120/2026. Encaminho para providências cabíveis.")
    saida = io.BytesIO()
    doc.save(saida)
    resultado, pdf, _ = analisar(saida.getvalue(), "despacho.docx")
    assert pdf.startswith(b"%PDF") and resultado["partes"][0]["tipo"] == "despacho"


def test_ocr_real_de_imagem():
    from PIL import Image, ImageDraw, ImageFont

    imagem = Image.new("RGB", (1500, 1000), "white")
    fonte = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 42)
    ImageDraw.Draw(imagem).multiline_text(
        (80, 80),
        "DESPACHO\nReferente ao Ofício 120/2026\nEncaminho para providências administrativas.",
        font=fonte,
        fill="black",
        spacing=35,
    )
    buffer = io.BytesIO()
    imagem.save(buffer, format="PNG")
    resultado, pdf, _ = analisar(buffer.getvalue(), "digitalizado.png")
    assert resultado["ocr"] == 1
    assert resultado["partes"][0]["tipo"] == "despacho"
    assert "120/2026" in resultado["partes"][0]["texto"] and pdf.startswith(b"%PDF")


def test_formato_desconhecido_e_pdf_invalido():
    for dados, nome in [(b"programa nao executavel", "arquivo.exe"), (b"%PDF-quebrado", "x.pdf")]:
        with pytest.raises(LeituraInvalida):
            analisar(dados, nome)


def test_ocr_de_corpo_digitalizado_com_rodape_textual():
    from PIL import Image, ImageDraw, ImageFont
    from weasyprint import HTML

    imagem = Image.new("RGB", (1500, 1000), "white")
    fonte = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 42)
    ImageDraw.Draw(imagem).multiline_text(
        (80, 80),
        "DESPACHO\nReferente ao Ofício 120/2026\nEncaminho para providências.",
        font=fonte,
        fill="black",
        spacing=35,
    )
    buffer = io.BytesIO()
    imagem.save(buffer, format="PNG")
    corpo = base64.b64encode(buffer.getvalue()).decode()
    dados = HTML(
        string=f'<img style="width:100%" src="data:image/png;base64,{corpo}">'
        "<p>Documento recebido pelo sistema eletrônico. Verificação da assinatura "
        "disponível no sistema. Este rodapé é texto, mas o corpo é digitalizado.</p>"
    ).write_pdf()
    resultado, _, _ = analisar(dados, "processo-digitalizado.pdf")
    assert resultado["ocr"] == 1
    assert resultado["partes"][0]["tipo"] == "despacho"
    assert "120/2026" in resultado["partes"][0]["texto"]


def test_processo_misto_anexa_nos_destinos_e_replay_nao_duplica(c):
    op = c.usuarios["operador"]
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    termo = termos.salvar(op, oficio=oficio)
    servidor = termos.efetivo(termo).servidores[0]
    numero = oficio.numero_formatado
    dados = pdf_textos(
        [
            f"Ofício {numero}\nSolicito autorização para a viagem da equipe ao evento.",
            f"DESPACHO\nReferente ao Ofício {numero}\nEncaminho para providências administrativas.",
            f"TERMO DE AUTORIZAÇÃO\nOfício {numero}\n{servidor.nome}\nCPF {servidor.cpf}\n"
            "Manifesto interesse na participação no evento.",
        ]
    )
    arquivo, novo = importacao.receber(op, "processo.pdf", dados)
    assert novo and importacao.reservar() == arquivo.pk
    importacao.processar(arquivo.pk)
    arquivo.refresh_from_db()
    assert arquivo.estado == "concluido", list(arquivo.partes.values("motivo", "estado"))
    assert list(arquivo.partes.values_list("estado", flat=True)) == ["anexado"] * 3
    assert ViaAssinada.objects.filter(oficio=oficio, tipo="oficio").count() == 1
    assert ViaAssinada.objects.filter(termo=termo, chave=str(servidor.pk)).count() == 1
    assert AnexoPrestacao.objects.filter(prestacao__oficio=oficio, tipo="despacho").count() == 1
    repetido, novo = importacao.receber(op, "renomeado.pdf", dados)
    assert not novo and repetido.pk == arquivo.pk
    importacao.processar(arquivo.pk)
    assert arquivo.partes.count() == 3 and ViaAssinada.objects.count() == 2


def test_revisao_privada_permissoes_e_substituicao(c):
    op = c.usuarios["operador"]
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    dados = pdf_textos(["Carta sem título reconhecível. Confira o destino antes de anexar."])
    arquivo, _ = importacao.receber(op, "duvida.pdf", dados)
    importacao.processar(arquivo.pk)
    parte = arquivo.partes.get()
    assert parte.estado == "revisao" and not parte.destino
    cliente = Client()
    cliente.force_login(c.usuarios["outra"])
    assert cliente.get(reverse("viagens:arquivo_importacao", args=[parte.pk])).status_code == 404
    assert cliente.get(reverse("viagens:original_importacao", args=[arquivo.pk])).status_code == 404
    cliente.force_login(c.usuarios["consulta"])
    assert cliente.get(reverse("viagens:importacoes")).status_code == 403
    with pytest.raises(PermissionDenied):
        importacao.receber(c.usuarios["consulta"], "teste.pdf", dados)
    cliente.force_login(op)
    assert cliente.get(reverse("viagens:revisar_importacao", args=[parte.pk])).status_code == 200
    resposta = cliente.post(
        reverse("viagens:revisar_importacao", args=[parte.pk]), {"destino": f"oficio:{oficio.pk}:"}
    )
    assert resposta.status_code == 302
    parte.refresh_from_db()
    assert parte.estado == "anexado" and parte.revisado_por == op
    novo, _ = importacao.receber(
        op,
        "nova.pdf",
        pdf_textos(
            [f"Ofício {oficio.numero_formatado}\nOutra via com conteúdo diferente para teste."]
        ),
    )
    importacao.processar(novo.pk)
    segunda = novo.partes.get()
    assert segunda.estado == "revisao" and "Já existe" in segunda.motivo
    importacao.aplicar(op, segunda.pk, f"oficio:{oficio.pk}:", confirmar=True)
    segunda.refresh_from_db()
    assert segunda.estado == "anexado" and segunda.revisado_por == op
    assert ViaAssinada.objects.filter(oficio=oficio, revogada_em__isnull=True).count() == 1


def test_lote_aceita_bons_e_recusa_invalido_sem_apagar_o_original(c):
    cliente = Client()
    cliente.force_login(c.usuarios["operador"])
    resposta = cliente.post(
        reverse("viagens:enviar_importacao"),
        {
            "arquivos": [
                SimpleUploadedFile(
                    "despacho.txt", b"DESPACHO\nArquivo com referencia ainda desconhecida."
                ),
                SimpleUploadedFile("desconhecido.xyz", b"formato desconhecido"),
            ]
        },
        HTTP_X_REQUESTED_WITH="fetch",
    )
    assert resposta.status_code == 202 and len(resposta.json()["arquivos"]) == 2
    for item in resposta.json()["arquivos"]:
        importacao.processar(item["id"])
    assert ArquivoImportado.objects.filter(estado="erro").count() == 1
    assert all(a.original.storage.exists(a.original.name) for a in ArquivoImportado.objects.all())
    erro = ArquivoImportado.objects.get(estado="erro")
    reenviado, novo = importacao.receber(
        c.usuarios["operador"], "outra.xyz", b"formato desconhecido"
    )
    assert not novo and reenviado.pk == erro.pk and reenviado.estado == "fila"


def test_upload_real_de_25_mib_passando_pelo_middleware(c):
    cliente = Client()
    cliente.force_login(c.usuarios["operador"])
    resposta = cliente.post(
        reverse("viagens:enviar_importacao"),
        {"arquivos": SimpleUploadedFile("grande.txt", b"x" * (25 * 1024**2))},
        HTTP_X_REQUESTED_WITH="fetch",
    )
    assert resposta.status_code == 202
    registro = ArquivoImportado.objects.get(pk=resposta.json()["arquivos"][0]["id"])
    assert registro.original.size == 25 * 1024**2 and registro.estado == "fila"


def test_lease_expirado_nao_bloqueia_fila_e_retorna_trabalho(c):
    op = c.usuarios["operador"]
    esgotado, _ = importacao.receber(op, "antigo.txt", b"antigo")
    retomado, _ = importacao.receber(op, "interrompido.txt", b"interrompido")
    novo, _ = importacao.receber(op, "novo.txt", b"novo")
    ArquivoImportado.objects.filter(pk__in=[esgotado.pk, retomado.pk]).update(
        estado="lendo", iniciado_em=timezone.now() - timedelta(minutes=10), tentativas=1
    )
    ArquivoImportado.objects.filter(pk=esgotado.pk).update(tentativas=3)
    assert importacao.reservar() == retomado.pk
    assert importacao.reservar() == novo.pk
    assert importacao.reservar() is None
    esgotado.refresh_from_db()
    retomado.refresh_from_db()
    assert esgotado.estado == "erro" and retomado.tentativas == 2


def test_documentos_reais_do_sistema_sem_referencia_artificial(c):
    op = c.usuarios["operador"]
    oficio = Oficio.objects.get(pk=c.ids["oficio_emitido"])
    documento = assinados.documento_emitido(oficio, "oficio")
    with documento.arquivo.open("rb") as f:
        arquivo, _ = importacao.receber(op, "oficio.pdf", f.read())
    importacao.processar(arquivo.pk)
    assert arquivo.partes.get().destino == f"oficio:{oficio.pk}:"
    assert arquivo.partes.get().estado == "anexado"
    termo = termos.salvar(op, oficio=oficio)
    servidor = termos.efetivo(termo).servidores[0]
    dados = termos.pdf_do_documento(termos.dados_do_documento(termo, str(servidor.pk)))
    arquivo, _ = importacao.receber(op, "termo.pdf", dados)
    importacao.processar(arquivo.pk)
    partes = list(arquivo.partes.all())
    assert len(partes) == 1
    assert partes[0].estado == "anexado", partes[0].motivo
    assert partes[0].destino == f"termo:{termo.pk}:{servidor.pk}"
