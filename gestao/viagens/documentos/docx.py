"""DOCX do documento a partir do mesmo HTML do PDF (D4; paridade com a referência).

A referência entrega "Baixar DOCX · Arquivo editável" do ofício e da justificativa, e o DOCX
da versão editada (m057) com conteúdo e formatação de texto, sem a geometria do PDF. Aqui é
igual: o HTML final (modelo + texto editado no editor + campos vivos) vira parágrafos com
alinhamento, negrito, itálico, sublinhado, listas, tabelas, o brasão, o rodapé e as quebras
de página. Larguras exatas, entrelinhas medidas e a fonte do PDF ficam de fora — quem
precisa do documento fiel é o PDF/A; o DOCX é para levar o texto e ajustá-lo fora do sistema.
"""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from urllib.parse import unquote, urlparse
from urllib.request import url2pathname

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Mm, Pt
from lxml import html as lxml_html

BLOCOS = {"p", "div", "h1", "h2", "h3", "h4", "h5", "h6", "li", "blockquote", "header",
          "section", "article", "main", "footer", "ul", "ol", "table"}
CENTRO = {"centro", "titulo-secao", "assinatura", "timbre"}
NEGRITO_CLASSES = {"titulo-secao", "destinatario", "nome", "cargo"}
IGNORAR = {"minuta", "quebra", "editor-marca", "sr-only"}
ALINHAMENTOS = {"left": WD_ALIGN_PARAGRAPH.LEFT, "right": WD_ALIGN_PARAGRAPH.RIGHT,
                "center": WD_ALIGN_PARAGRAPH.CENTER, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


def _classes(el) -> set[str]:
    return set((el.get("class") or "").split())


def _herda(el, conjunto: set[str]) -> bool:
    """O elemento ou algum ancestral tem uma das classes."""
    while el is not None:
        if _classes(el) & conjunto or (el.tag == "header" and "timbre" in conjunto):
            return True
        el = el.getparent()
    return False


def _alinhamento(el):
    while el is not None:
        estilo = (el.get("style") or "").replace(" ", "")
        for nome, valor in ALINHAMENTOS.items():
            if f"text-align:{nome}" in estilo:
                return valor
        if _classes(el) & CENTRO or el.tag in ("th", "header"):
            return WD_ALIGN_PARAGRAPH.CENTER
        el = el.getparent()
    return None


def _tem_bloco(el) -> bool:
    return any(isinstance(f.tag, str) and f.tag in BLOCOS for f in el.iterchildren())


class _Escritor:
    def __init__(self, base_imagens: Path):
        self.doc = Docx()
        secao = self.doc.sections[0]
        secao.page_height, secao.page_width = Mm(297), Mm(210)
        secao.top_margin, secao.bottom_margin = Mm(14), Mm(24)
        secao.left_margin = secao.right_margin = Mm(18)
        estilo = self.doc.styles["Normal"]
        estilo.font.name = "Times New Roman"
        estilo.font.size = Pt(11)
        self.base_imagens = base_imagens
        self.quebra_pendente = False

    # ------------------------------------------------------------ trechos (inline)
    def _trechos(self, paragrafo, el, *, negrito=False, italico=False, sublinhado=False):
        if el.text:
            self._run(paragrafo, el.text, negrito, italico, sublinhado)
        for filho in el.iterchildren():
            if not isinstance(filho.tag, str) or _classes(filho) & IGNORAR:
                if filho.tail:
                    self._run(paragrafo, filho.tail, negrito, italico, sublinhado)
                continue
            tag = filho.tag
            if tag == "br":
                paragrafo.add_run().add_break()
            elif tag == "img":
                self._imagem(paragrafo, filho)
            elif "rodape" in _classes(el) and tag == "strong":
                # No PDF a sigla do rodapé é um bloco acima do endereço.
                self._trechos(paragrafo, filho, negrito=True)
                paragrafo.add_run().add_break()
            else:
                self._trechos(
                    paragrafo, filho,
                    negrito=negrito or tag in ("b", "strong", "th")
                    or bool(_classes(filho) & NEGRITO_CLASSES),
                    italico=italico or tag in ("i", "em"),
                    sublinhado=sublinhado or tag == "u")
            if filho.tail:
                self._run(paragrafo, filho.tail, negrito, italico, sublinhado)

    @staticmethod
    def _run(paragrafo, texto, negrito, italico, sublinhado):
        """Espaços como o navegador: sequências viram um espaço, e o das pontas fica (é ele
        que separa "Delegado," de "através" quando há uma tag no meio)."""
        if not texto:
            return
        miolo = " ".join(texto.split())
        if not miolo:
            if paragrafo.runs and not paragrafo.runs[-1].text.endswith(" "):
                paragrafo.add_run(" ")
            return
        texto = (" " if texto[0].isspace() and paragrafo.runs else "") + miolo + (
            " " if texto[-1].isspace() else "")
        run = paragrafo.add_run(texto)
        run.bold, run.italic, run.underline = negrito or None, italico or None, sublinhado or None

    def _imagem(self, paragrafo, el):
        endereco = urlparse(el.get("src") or "")
        caminho = (Path(url2pathname(endereco.path)) if endereco.scheme == "file"
                   else self.base_imagens / Path(unquote(endereco.path)).name)
        if not caminho.is_file():
            caminho = self.base_imagens / caminho.name
        if caminho.is_file() and caminho.suffix.lower() in (".png", ".jpg", ".jpeg"):
            paragrafo.add_run().add_picture(str(caminho), height=Mm(18))

    # ------------------------------------------------------------ blocos
    def paragrafo(self, el, estilo: str | None = None):
        p = self.doc.add_paragraph(style=estilo) if estilo else self.doc.add_paragraph()
        if self.quebra_pendente:
            p.paragraph_format.page_break_before = True
            self.quebra_pendente = False
        alinhamento = _alinhamento(el)
        if alinhamento is not None:
            p.alignment = alinhamento
        negrito = el.tag in ("h1", "h2", "h3", "h4") or bool(_classes(el) & NEGRITO_CLASSES) \
            or _herda(el, {"titulo-secao", "destinatario"}) or el.tag == "header" \
            or (el.getparent() is not None and el.getparent().tag == "header")
        self._trechos(p, el, negrito=negrito)
        return p

    def tabela(self, el):
        linhas = list(el.iter("tr"))
        if not linhas:
            return
        colunas = max(sum(int(c.get("colspan") or 1) for c in tr if c.tag in ("td", "th"))
                      for tr in linhas)
        borda = "sem-borda" not in _classes(el)
        tabela = self.doc.add_table(rows=len(linhas), cols=max(1, colunas))
        if borda:
            tabela.style = "Table Grid"
        for i, tr in enumerate(linhas):
            j = 0
            for celula in (c for c in tr if c.tag in ("td", "th")):
                alvo = tabela.cell(i, j)
                largura = int(celula.get("colspan") or 1)
                if largura > 1:
                    alvo = alvo.merge(tabela.cell(i, min(colunas - 1, j + largura - 1)))
                paragrafo = alvo.paragraphs[0]
                alinhamento = _alinhamento(celula)
                if alinhamento is not None:
                    paragrafo.alignment = alinhamento
                self._trechos(paragrafo, celula, negrito=celula.tag == "th")
                j += largura

    def bloco(self, el):
        if not isinstance(el.tag, str) or _classes(el) & IGNORAR - {"quebra"}:
            return
        classes = _classes(el)
        if "quebra--ativa" in classes:
            self.quebra_pendente = True
            return
        if "quebra" in classes or el.tag in ("style", "script", "head", "title", "meta"):
            return
        if "rodape" in classes:
            rodape = self.doc.sections[0].footer.paragraphs[0]
            rodape.alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._trechos(rodape, el, italico=True)
            return
        if el.tag == "table":
            self.tabela(el)
        elif el.tag in ("ul", "ol"):
            estilo = "List Number" if el.tag == "ol" else "List Bullet"
            for li in el.iterchildren("li"):
                self.paragrafo(li, estilo)
        elif _tem_bloco(el):
            if (el.text or "").strip():
                self._run(self.doc.add_paragraph(), el.text, False, False, False)
            for filho in el.iterchildren():
                self.bloco(filho)
                if filho.tail and filho.tail.strip():
                    self._run(self.doc.add_paragraph(), filho.tail, False, False, False)
        elif el.tag == "img":
            self._imagem(self.doc.add_paragraph(), el)
        elif (el.text_content() or "").strip() or el.find(".//img") is not None:
            self.paragrafo(el)


def docx_do_html(html: str, *, base_imagens: Path, minuta: bool = False) -> bytes:
    """O HTML final do documento como DOCX. Minuta (rascunho) leva o aviso no cabeçalho."""
    arvore = lxml_html.document_fromstring(html)
    escritor = _Escritor(base_imagens)
    if minuta:
        aviso = escritor.doc.sections[0].header.paragraphs[0]
        aviso.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = aviso.add_run("MINUTA — documento ainda não emitido")
        run.bold = True
    corpo = arvore.find("body")
    for filho in (corpo if corpo is not None else arvore).iterchildren():
        escritor.bloco(filho)
    if escritor.quebra_pendente:
        escritor.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    saida = BytesIO()
    escritor.doc.save(saida)
    return saida.getvalue()
