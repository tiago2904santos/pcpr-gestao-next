"""O relatório consolidado (R1; paridade com `relatorios/consolidacao.py` da referência — a
planilha "PCPR na Comunidade - ASCOM" que juntava as outras): os números de todos os
módulos numa página, por mês de um ano (ou por ano, em "Todos os anos").

Duas regras, as mesmas da agenda: **a permissão é a de cada módulo** (uma seção fora do
acesso de quem pede não é calculada nem aparece; ofícios, viagens e solicitações de evento
passam pela visibilidade do próprio módulo); e **conta o que aconteceu** (palestras e
eventos da ASCOM atendidos; solicitações de evento atendidas, ou deferidas com a data já
passada; coffee break, viagens e ofícios não cancelados)."""

from __future__ import annotations

from collections import defaultdict
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import date

from django.db.models import Q

MESES = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
PRIMEIRO_ANO = 2022
ANO_MINIMO = 2000


@dataclass(frozen=True)
class Periodo:
    """Um ano (colunas por mês) ou todos os anos (`ano=None`, colunas por ano)."""

    ano: int | None
    hoje: date

    @property
    def rotulo(self) -> str:
        return str(self.ano) if self.ano else "Todos os anos"

    @property
    def coluna(self) -> str:
        return "Mês/ano" if self.ano else "Ano"

    def chaves(self, anos_com_dados: Iterable[int] = ()) -> list:
        if self.ano:
            return [(self.ano, m) for m in range(1, 13)]
        return sorted(set(anos_com_dados) | {self.hoje.year})

    def chave(self, d: date):
        return (d.year, d.month) if self.ano else d.year

    def rotulo_da_chave(self, chave) -> str:
        return f"{MESES[chave[1] - 1]}/{str(chave[0])[2:]}" if self.ano else str(chave)

    def contem(self, d: date | None) -> bool:
        if d is None:
            return False
        if self.ano:
            return d.year == self.ano
        return ANO_MINIMO <= d.year <= self.hoje.year + 1  # datas digitadas erradas ficam fora

    def filtrar(self, qs, campo: str):
        if self.ano:
            return qs.filter(**{f"{campo}__year": self.ano})
        return qs.filter(**{f"{campo}__year__gte": ANO_MINIMO,
                            f"{campo}__year__lte": self.hoje.year + 1})


def anos_disponiveis(hoje: date) -> list[int]:
    return list(range(hoje.year, PRIMEIRO_ANO - 1, -1))


@dataclass
class Secao:
    slug: str
    titulo: str
    icone: str
    colunas: list[str]
    linhas: list[list]
    totais: list = field(default_factory=list)
    nota: str = ""
    numericas: frozenset[int] = frozenset()


@dataclass
class Indicador:
    titulo: str
    valor: object
    icone: str
    detalhe: str


@dataclass
class Relatorio:
    periodo: Periodo
    indicadores: list[Indicador] = field(default_factory=list)
    quadro: Secao | None = None
    secoes: list[Secao] = field(default_factory=list)


def _por_mes(periodo: Periodo, itens: Iterable, data_de: Callable, valores_de: Callable,
             n: int) -> tuple[list[list], list]:
    """Soma os itens por mês (ou por ano) em `n` colunas, com a linha de total."""
    baldes: dict = defaultdict(lambda: [0] * n)
    for item in itens:
        d = data_de(item)
        if periodo.contem(d):
            for i, v in enumerate(valores_de(item)):
                baldes[periodo.chave(d)][i] += v or 0
    chaves = periodo.chaves(set(baldes) if not periodo.ano else ())
    linhas = [[periodo.rotulo_da_chave(c), *baldes.get(c, [0] * n)] for c in chaves]
    return linhas, ["Total", *[sum(linha[i + 1] for linha in linhas) for i in range(n)]]


def _data_br(d: date | None) -> str:
    return f"{d:%d/%m/%Y}" if d else ""


def _mes_de(obj) -> date | None:
    """O "mês" da planilha: o do evento; sem data de evento, o da solicitação."""
    return obj.data_inicio_evento or obj.data_solicitacao


# ---------------------------------------------------------------- fontes
def _palestras(usuario, periodo: Periodo, evento: str) -> list:
    from gestao.palestras import policies
    from gestao.palestras.models import Palestra

    if not policies.pode_acessar(usuario):
        return []
    qs = (Palestra.objects.filter(evento=evento, status=Palestra.Status.ATENDIDA)
          .select_related("municipio").prefetch_related("temas"))
    if periodo.ano:
        qs = qs.filter(Q(data_inicio_evento__year=periodo.ano)
                       | Q(data_inicio_evento__isnull=True, data_solicitacao__year=periodo.ano))
    return list(qs)


def _solicitacoes_realizadas(usuario, periodo: Periodo) -> list:
    from gestao.eventos import dominio, policies

    qs = policies.solicitacoes_visiveis(usuario).filter(
        Q(status=dominio.ATENDIDA)
        | Q(status=dominio.DEFERIDA, data_inicio_evento__lte=periodo.hoje))
    return list(periodo.filtrar(qs, "data_inicio_evento")
                .select_related("municipio", "tipo_evento").order_by("data_inicio_evento", "pk"))


def _eh_pcpr(s) -> bool:
    return "COMUNIDADE" in (s.tipo_evento.nome if s.tipo_evento else "").upper()


def _cidade(obj) -> str:
    return obj.municipio.nome if getattr(obj, "municipio", None) else ""


# ---------------------------------------------------------------- seções
def secao_palestras(usuario, periodo: Periodo) -> Secao:
    linhas, totais = _por_mes(periodo, _palestras(usuario, periodo, "palestra"), _mes_de,
                              lambda p: (1, p.quantidade_publico), 2)
    return Secao("palestras", "Palestras", "presentation",
                 [periodo.coluna, "Qtd.", "Pessoas atendidas"], linhas, totais,
                 "Palestras com status Atendida no módulo Palestras e eventos, no mês do evento.",
                 frozenset({1, 2}))


def secao_pcpr(usuario, periodo: Periodo, solicitacoes: list) -> tuple[Secao, list[dict]]:
    """As edições do PCPR na Comunidade das duas fontes, sem contar a mesma duas vezes: a
    solicitação de evento (que traz as CIN) e a palestra atendida (que traz o público)
    viram uma linha só quando falam do mesmo município no período do evento."""
    edicoes: list[dict] = [
        {"data": s.data_inicio_evento, "fim": s.data_fim_evento or s.data_inicio_evento,
         "municipio_id": s.municipio_id, "cidade": _cidade(s), "publico": None,
         "cin": s.quantidade_cin, "origem": "Solicitação"}
        for s in solicitacoes if _eh_pcpr(s)]
    for p in _palestras(usuario, periodo, "pcpr_na_comunidade"):
        d = p.data_inicio_evento
        par = next((e for e in edicoes if e["origem"] == "Solicitação" and e["publico"] is None
                    and p.municipio_id and e["municipio_id"] == p.municipio_id and d
                    and e["data"] <= d <= e["fim"]), None)
        if par is not None:
            par["publico"], par["origem"] = p.quantidade_publico, "Solicitação + ASCOM"
            continue
        edicoes.append({"data": _mes_de(p), "fim": _mes_de(p), "municipio_id": p.municipio_id,
                        "cidade": _cidade(p), "publico": p.quantidade_publico, "cin": None,
                        "origem": "ASCOM"})
    edicoes.sort(key=lambda e: (e["data"], e["cidade"]))
    tabela = [[f"{MESES[e['data'].month - 1]}/{e['data']:%y}", _data_br(e["data"]),
               (e["cidade"] or "Sem município").upper(),
               "" if e["publico"] is None else e["publico"], "" if e["cin"] is None else e["cin"],
               e["origem"]] for e in edicoes]
    n = len(edicoes)
    totais = [f"{n} ediç{'ões' if n != 1 else 'ão'}", "", "",
              sum(e["publico"] or 0 for e in edicoes), sum(e["cin"] or 0 for e in edicoes), ""]
    return Secao("pcpr", "PCPR na Comunidade", "landmark",
                 ["Mês/ano", "Data", "Cidade", "Público", "Qtd. RGs (CIN)", "Origem"], tabela,
                 totais, "Solicitações de evento do tipo PCPR na Comunidade (atendidas ou "
                 "deferidas com data já passada), somadas às palestras atendidas desse tipo. A "
                 "mesma edição nas duas fontes vira uma linha só.", frozenset({3, 4})), edicoes


def secao_eventos(usuario, periodo: Periodo, solicitacoes: list) -> tuple[Secao, list[tuple]]:
    itens = []
    for s in solicitacoes:
        if _eh_pcpr(s):
            continue
        descricao = s.local_evento or (s.descricao_complementar or "").strip().split("\n")[0]
        itens.append((s.data_inicio_evento, s.tipo_evento.nome if s.tipo_evento else "Evento",
                      descricao, _cidade(s), "Solicitação"))
    for p in _palestras(usuario, periodo, "evento"):
        descricao = (p.assunto_email or (p.descricao or "").strip().split("\n")[0]
                     or ", ".join(t.nome for t in p.temas.all()))
        itens.append((_mes_de(p), "Evento (ASCOM)", descricao, _cidade(p), "ASCOM"))
    itens.sort(key=lambda i: (i[0], i[1]))
    n = len(itens)
    return Secao("eventos", "Eventos em geral", "calendar",
                 ["Data", "Tipo", "Descrição", "Local", "Origem"],
                 [[_data_br(d), tipo, desc[:160], local, origem]
                  for d, tipo, desc, local, origem in itens],
                 [f"{n} evento{'s' if n != 1 else ''}", "", "", "", ""],
                 "Solicitações de evento dos demais tipos (inaugurações, solenidades, ações…) "
                 "e os eventos atendidos do módulo Palestras e eventos."), itens


def secao_coffee(usuario, periodo: Periodo) -> Secao:
    from gestao.coffee.models import Solicitacao

    qs = Solicitacao.objects.all()
    if periodo.ano:
        qs = qs.filter(Q(data_evento__year=periodo.ano)
                       | Q(data_evento__isnull=True, data_solicitacao__year=periodo.ano))
    linhas, totais = _por_mes(
        periodo, qs.only("data_evento", "data_solicitacao", "cancelada", "quantidade",
                         "quantidade_faturada", "valor_unitario"),
        lambda s: s.data_evento or s.data_solicitacao,
        lambda s: (0, 0, 1, 0) if s.cancelada else (1, s.quantidade_efetiva, 0, s.valor), 4)
    return Secao("coffee", "Coffee break", "receipt",
                 [periodo.coluna, "Solicitações", "Quantidade servida", "Canceladas",
                  "Valor (R$)"], linhas, totais,
                 "Solicitações de coffee break no mês do evento; a quantidade e o valor "
                 "(quantidade × preço unitário guardado na OS) somam só as não canceladas.",
                 frozenset({1, 2, 3, 4}))


def secao_publicacoes(usuario, periodo: Periodo) -> Secao:
    from gestao.publicacoes.models import Publicacao

    itens = periodo.filtrar(Publicacao.objects.all(), "data").only("data", "status")
    linhas, totais = _por_mes(periodo, itens, lambda p: p.data,
                              lambda p: (1, int(p.status == Publicacao.Status.PUBLICADA)), 2)
    return Secao("publicacoes", "Publicações", "newspaper",
                 [periodo.coluna, "Pautas", "Publicadas"], linhas, totais,
                 "Pautas pela data da pauta.", frozenset({1, 2}))


def secao_imprensa(usuario, periodo: Periodo) -> Secao:
    from gestao.imprensa.models import Atendimento

    itens = periodo.filtrar(Atendimento.objects.all(), "data").only("data", "situacao")
    linhas, totais = _por_mes(periodo, itens, lambda a: a.data,
                              lambda a: (1, int(a.situacao == Atendimento.Situacao.ATENDIDO)), 2)
    return Secao("imprensa", "Atendimento à imprensa", "megaphone",
                 [periodo.coluna, "Pedidos", "Atendidos"], linhas, totais,
                 "Pedidos de jornalistas pela data do pedido.", frozenset({1, 2}))


def secao_viagens(usuario, periodo: Periodo) -> Secao:
    from gestao.viagens import policies
    from gestao.viagens.models import Oficio, Viagem

    viagens = periodo.filtrar(policies.viagens_visiveis(usuario)
                              .exclude(situacao=Viagem.Situacao.CANCELADA), "data_inicio")
    oficios = periodo.filtrar(policies.oficios_visiveis(usuario)
                              .exclude(situacao=Oficio.Situacao.CANCELADO), "criado_em")
    itens = ([("v", v) for v in viagens.values_list("data_inicio", flat=True)]
             + [("o", o.date()) for o in oficios.values_list("criado_em", flat=True)])
    linhas, totais = _por_mes(periodo, itens, lambda i: i[1],
                              lambda i: (int(i[0] == "v"), int(i[0] == "o")), 2)
    return Secao("viagens", "Viagens", "plane", [periodo.coluna, "Viagens", "Ofícios"],
                 linhas, totais, "Viagens pela data de saída e ofícios pela data de criação, "
                 "sem os cancelados (só os que você vê).", frozenset({1, 2}))


# ---------------------------------------------------------------- o relatório
def _pode(modulo: str, usuario) -> bool:
    if modulo == "palestras":
        from gestao.palestras import policies as p
        return p.pode_acessar(usuario)
    if modulo == "coffee":
        from gestao.coffee import policies as c
        return c.pode_acessar(usuario)
    if modulo == "publicacoes":
        from gestao.publicacoes import policies as pu
        return pu.pode_acessar(usuario)
    if modulo == "imprensa":
        from gestao.imprensa import policies as i
        return i.pode_acessar(usuario)
    if modulo == "viagens":
        from gestao.viagens import policies as v
        return v.pode_listar(usuario)
    return False


def montar(usuario, ano: int | None, hoje: date) -> Relatorio:
    periodo = Periodo(ano, hoje)
    r = Relatorio(periodo)
    visao: list[tuple[str, dict]] = []

    def por_chave(datas: Iterable[date | None]) -> dict:
        contagem: dict = defaultdict(int)
        for d in datas:
            if d is not None and periodo.contem(d):
                contagem[periodo.chave(d)] += 1
        return contagem

    def coluna(secao: Secao, i: int) -> dict:
        if periodo.ano:
            return dict(zip(periodo.chaves(), [linha[i] for linha in secao.linhas],
                            strict=True))
        return {int(linha[0]): linha[i] for linha in secao.linhas}

    if _pode("palestras", usuario):
        s = secao_palestras(usuario, periodo)
        r.secoes.append(s)
        r.indicadores.append(Indicador("Palestras", s.totais[1], "presentation",
                                       f"{s.totais[2]} pessoas atendidas"))
        visao.append(("Palestras", coluna(s, 1)))
    solicitacoes = _solicitacoes_realizadas(usuario, periodo)
    pcpr, edicoes = secao_pcpr(usuario, periodo, solicitacoes)
    r.secoes.append(pcpr)
    r.indicadores.append(Indicador("PCPR na Comunidade", len(edicoes), "landmark",
                                   f"{pcpr.totais[4]} CIN emitidas"))
    visao.append(("PCPR na Comunidade", por_chave(e["data"] for e in edicoes)))
    eventos, itens = secao_eventos(usuario, periodo, solicitacoes)
    r.secoes.append(eventos)
    r.indicadores.append(Indicador("Eventos em geral", len(itens), "calendar",
                                   "inaugurações, solenidades e ações"))
    visao.append(("Eventos em geral", por_chave(i[0] for i in itens)))
    mensais = (("coffee", secao_coffee, "Coffee break", 1, 2, "unidades servidas"),
               ("publicacoes", secao_publicacoes, "Publicações", 2, 1, "pautas"),
               ("imprensa", secao_imprensa, "Atendimento à imprensa", 1, 2, "atendidos"),
               ("viagens", secao_viagens, "Viagens", 1, 2, "ofícios"))
    for modulo, fazer, rotulo, col_kpi, col_detalhe, detalhe in mensais:
        if not _pode(modulo, usuario):
            continue
        s = fazer(usuario, periodo)
        r.secoes.append(s)
        r.indicadores.append(Indicador(rotulo, s.totais[col_kpi], s.icone,
                                       f"{s.totais[col_detalhe]} {detalhe}"))
        visao.append((rotulo, coluna(s, col_kpi)))
    anos = {c for _r, contagem in visao for c in contagem if isinstance(c, int)}
    chaves = periodo.chaves(anos)
    r.quadro = Secao(
        "geral", "Visão geral por mês" if periodo.ano else "Visão geral por ano", "table",
        ["Módulo", *[periodo.rotulo_da_chave(c) for c in chaves], "Total"],
        [[rotulo, *[contagem.get(c, 0) for c in chaves], sum(contagem.get(c, 0) for c in chaves)]
         for rotulo, contagem in visao],
        nota="Palestras, PCPR e eventos contam o que aconteceu; coffee break, solicitações não "
             "canceladas; publicações, pautas publicadas; imprensa, pedidos recebidos; viagens, "
             "viagens não canceladas.",
        numericas=frozenset(range(1, len(chaves) + 2)))
    return r


def _seguro(valor):
    """Texto que começaria uma fórmula no Excel vai como texto literal."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + valor
    return valor


def planilha(r: Relatorio) -> bytes:
    """O mesmo relatório em XLSX: uma aba por seção (a visão geral primeiro)."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    if livro.active is not None:
        livro.remove(livro.active)
    for s in [r.quadro, *r.secoes]:
        if s is None:
            continue
        aba = livro.create_sheet(s.titulo[:31])
        aba.append([f"{s.titulo.upper()} — {r.periodo.rotulo}"])
        aba["A1"].font = Font(bold=True, size=13)
        aba.append(s.colunas)
        for celula in aba[2]:
            celula.font = Font(bold=True)
        for linha in s.linhas:
            aba.append([_seguro(v) for v in linha])
        if s.totais:
            aba.append([_seguro(v) for v in s.totais])
            for celula in aba[aba.max_row]:
                celula.font = Font(bold=True)
        if s.nota:
            aba.append([])
            aba.append([s.nota])
            aba.cell(aba.max_row, 1).font = Font(italic=True)
        for i, coluna in enumerate(s.colunas, start=1):
            maior = max([len(str(coluna))] + [len(str(linha[i - 1])) for linha in s.linhas
                                               if i - 1 < len(linha)])
            aba.column_dimensions[get_column_letter(i)].width = min(max(10, maior + 2), 60)
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
