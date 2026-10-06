"""R1: o relatório consolidado — período por mês/ano, as seções de cada módulo só para quem
tem o módulo, o PCPR na Comunidade sem contar duas vezes a mesma edição, o coffee break
sem as canceladas no valor, a tela e a planilha com os mesmos números."""

from __future__ import annotations

import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from django.contrib.auth.models import Group
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from gestao.cadastros.models import Municipio
from gestao.identidade.models import Usuario
from gestao.painel import consolidado

pytestmark = pytest.mark.django_db


def test_periodo():
    p = consolidado.Periodo(2026, date(2026, 10, 6))
    assert p.chaves()[0] == (2026, 1) and len(p.chaves()) == 12
    assert p.rotulo_da_chave((2026, 3)) == "mar/26" and p.coluna == "Mês/ano"
    todos = consolidado.Periodo(None, date(2026, 10, 6))
    assert todos.rotulo == "Todos os anos" and todos.chaves({2024}) == [2024, 2026]
    assert not todos.contem(date(1905, 1, 1)) and todos.contem(date(2024, 5, 1))


@pytest.fixture
def cenario():
    hoje = timezone.localdate()
    ano = hoje.year
    cwb = Municipio.objects.get_or_create(codigo_ibge="4106902",
                                          defaults={"nome": "Curitiba", "uf": "PR"})[0]
    todos = Usuario.objects.create_user("chefe", "chefe@teste.invalid", None, nome="Chefe")
    for papel in ("ASCOM_PALESTRAS", "ASCOM_COFFEE_BREAK", "ASCOM_IMPRENSA", "GESTOR_DG"):
        todos.groups.add(Group.objects.get(name=papel))
    from gestao.eventos import dominio
    from gestao.eventos.models import Solicitacao, TipoEvento
    from gestao.palestras.models import Palestra

    dia = hoje  # o mesmo ano sempre (a palestra do PCPR cai no período da solicitação)
    pcpr = Solicitacao.objects.create(
        data_solicitacao=dia, data_inicio_evento=dia, data_fim_evento=dia + timedelta(days=1),
        municipio=cwb, tipo_evento=TipoEvento.objects.get(nome="PCPR na Comunidade"),
        quantidade_cin=120, status=dominio.ATENDIDA, criado_por=todos)
    Solicitacao.objects.create(data_solicitacao=dia, data_inicio_evento=dia, municipio=cwb,
                               tipo_evento=TipoEvento.objects.get(nome="Feira"),
                               local_evento="Ginásio", status=dominio.ATENDIDA,
                               criado_por=todos)
    for evento, publico, d in (("palestra", 80, dia), ("pcpr_na_comunidade", 300, dia)):
        Palestra.objects.create(data_solicitacao=dia, solicitante="Escola", evento=evento,
                                status=Palestra.Status.ATENDIDA, data_inicio_evento=d,
                                municipio=cwb, quantidade_publico=publico, criado_por=todos)
    from gestao.coffee.models import Contrato, Fornecedor, Lote
    from gestao.coffee.models import Solicitacao as Ordem

    f = Fornecedor.objects.create(razao_social="Buffet Exemplo Ltda", cnpj="11222333000181")
    c = Contrato.objects.create(fornecedor=f, numero="1/2026", valor_unitario=Decimal("20"))
    lote = Lote.objects.create(contrato=c, numero=1, exercicio=str(ano), quantidade_total=500)
    for qtd, cancelada in ((40, False), (10, True)):
        Ordem.objects.create(lote=lote, municipio=cwb, data_solicitacao=dia, data_evento=dia,
                             descricao="Evento", quantidade=qtd, valor_unitario=Decimal("20"),
                             cancelada=cancelada, criado_por=todos)
    return todos, pcpr, dia


def test_secoes_e_numeros(cenario):
    todos, _pcpr, dia = cenario
    r = consolidado.montar(todos, dia.year, timezone.localdate())
    slugs = [s.slug for s in r.secoes]
    assert slugs[:3] == ["palestras", "pcpr", "eventos"] and "coffee" in slugs
    pcpr_secao = next(s for s in r.secoes if s.slug == "pcpr")
    assert len(pcpr_secao.linhas) == 1  # solicitação + palestra do mesmo município = uma
    assert pcpr_secao.linhas[0][3:] == [300, 120, "Solicitação + ASCOM"]
    eventos = next(s for s in r.secoes if s.slug == "eventos")
    assert [linha[1:4] for linha in eventos.linhas] == [["Feira", "Ginásio", "Curitiba"]]
    coffee = next(s for s in r.secoes if s.slug == "coffee")
    assert coffee.totais == ["Total", 1, 40, 1, Decimal("800.00")]
    kpis = {k.titulo: k.valor for k in r.indicadores}
    assert kpis["Palestras"] == 1 and kpis["PCPR na Comunidade"] == 1
    assert kpis["Coffee break"] == 1
    quadro = {linha[0]: linha[-1] for linha in r.quadro.linhas}
    assert quadro["Eventos em geral"] == 1 and quadro["Coffee break"] == 1


def test_secao_so_para_quem_tem_o_modulo(cenario):
    sem_modulo = Usuario.objects.create_user("zeca", "z@teste.invalid", None, nome="Zeca")
    r = consolidado.montar(sem_modulo, None, timezone.localdate())
    assert [s.slug for s in r.secoes] == ["pcpr", "eventos"]  # só as próprias solicitações
    assert r.secoes[0].linhas == [] and r.secoes[1].linhas == []


def test_tela_e_planilha(cenario):
    todos, _pcpr, dia = cenario
    cli = Client()
    cli.force_login(todos)
    html = cli.get(reverse("painel:relatorios") + f"?ano={dia.year}").content.decode()
    assert "Relatório consolidado" in html and "PCPR na Comunidade" in html
    assert 'aria-current="page">' + str(dia.year) in html
    r = cli.get(reverse("painel:exportar_relatorio") + f"?ano={dia.year}")
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    from openpyxl import load_workbook
    livro = load_workbook(io.BytesIO(r.content))
    assert livro.sheetnames[0] == "Visão geral por mês" and "Coffee break" in livro.sheetnames
    assert cli.get(reverse("painel:relatorios") + "?ano=1999").status_code == 200  # vira o atual
