"""Telas das solicitações de evento (paridade com `solicitacoes/views.py` da referência):
lista com as filas, a folha (nova, rascunho que se grava sozinho, leitura depois do envio,
reabrir para alterar), as ações (enviar, despacho da DG com "registrar e abrir a próxima",
ajuste de servidores, concluir, cancelar, transferir, duplicar, excluir), anexos e a
exportação."""

from __future__ import annotations

import csv
import re
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import F
from django.http import FileResponse, Http404, HttpRequest, HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.http import require_GET, require_http_methods, require_POST

from gestao.cadastros.models import Municipio

from . import conflitos, dominio, ganchos, pedir_coffee, policies, queries, solicitacoes
from .forms import FORM_ID, FormularioSolicitacao, estrutura_do_post, quantidade
from .models import AnexoSolicitacao, Equipe, Servico, Solicitacao, TextoDespacho, TipoEvento

POR_PAGINA = 25


def _migalhas(*itens: tuple[str, str]) -> list[tuple[str, str]]:
    return [("Início", reverse("painel:inicio")),
            ("Solicitações de evento", reverse("eventos:solicitacoes")), *itens]


def _querystring(request: HttpRequest, *sem: str) -> str:
    return urlencode([(k, v) for k, vs in request.GET.lists() for v in vs if k not in sem and v])


def _visivel(request: HttpRequest, pk: int) -> Solicitacao:
    s = get_object_or_404(queries.base(request.user).prefetch_related(
        "servicos__servico", "equipes__equipe", "anexos"), pk=pk)
    return s


# ---------------------------------------------------------------- lista
@require_GET
def lista(request: HttpRequest) -> HttpResponse:
    if not policies.pode_criar(request.user):
        raise PermissionDenied
    hoje = timezone.localdate()
    f = queries.ler_filtros(request.GET)
    sem_fila = queries.filtrar(request.user, f, hoje, com_fila=False)
    filas = [(c, r, queries.aplicar_fila(sem_fila, c, request.user, hoje).count())
             for c, r in queries.filas_visiveis(request.user)]
    qs = queries.filtrar(request.user, f, hoje)
    # A fila do despacho na ordem do "Registrar e abrir a próxima": o evento mais próximo.
    qs = (qs.order_by(F("data_inicio_evento").asc(nulls_last=True), "pk") if f.fila == "despacho"
          else qs.order_by("-data_solicitacao", "-criado_em"))
    pagina = Paginator(qs, POR_PAGINA).get_page(request.GET.get("pagina"))
    base = _querystring(request, "pagina")
    return render(request, "eventos/solicitacoes/lista.html", {
        "page_obj": pagina, "linhas": [queries.linha(s, hoje) for s in pagina.object_list],
        "filtros": f, "filas": filas, "qs_sem_fila": _querystring(request, "fila", "pagina"),
        "querystring_base": f"{base}&" if base else "",
        "municipios": Municipio.objects.filter(
            pk__in=policies.solicitacoes_visiveis(request.user).values("municipio_id")
        ).order_by("nome"),
        "tipos": TipoEvento.objects.order_by("nome"),
        "url_exportar": reverse("eventos:exportar") + (f"?{base}" if base else ""),
        "url_exportar_xlsx": (reverse("eventos:exportar")
                              + f"?{base}{'&' if base else ''}formato=xlsx"),
        "pode_despachar": policies.pode_despachar(request.user),
        "migalhas": [("Início", reverse("painel:inicio")), ("Solicitações de evento", "")],
    })


COLUNAS = ("Nº", "Status", "Data da solicitação", "Início do evento", "Fim do evento",
           "Município", "Tipo de evento", "Local", "Endereço", "Bairro", "CEP", "Protocolo",
           "Solicitante", "Cargo / unidade", "Contato", "Órgão responsável", "Serviços",
           "Equipes (servidores)", "Total de servidores", "Tipo de operação",
           "Unidade móvel", "Qtde CIN", "Motorista", "Decisão DG", "Observações DG",
           "Decidido por", "Decidido em", "Criado por")


def _d(valor) -> str:
    return f"{valor:%d/%m/%Y}" if valor else ""


def _linha_exportada(s) -> list:
    return [_celula(v) for v in (
        s.pk, s.get_status_display(), _d(s.data_solicitacao), _d(s.data_inicio_evento),
        _d(s.data_fim_evento), s.municipio.nome if s.municipio else "",
        s.tipo_evento.nome if s.tipo_evento else "", s.local_evento,
        s.endereco, s.bairro, s.cep, s.protocolo, s.solicitante_nome,
        s.solicitante_cargo_unidade, s.contato,
        s.orgao_responsavel.nome if s.orgao_responsavel else "",
        "; ".join(x.servico.nome for x in s.servicos.all()),
        "; ".join(f"{x.equipe.nome} ({x.quantidade_servidores or 0})" for x in s.equipes.all()),
        s.quantidade_servidores, s.get_tipo_operacao_display(),
        "Sim" if s.unidade_movel else "Não", s.quantidade_cin or "",
        s.motorista.nome if s.motorista else "", s.get_decisao_dg_display(),
        s.observacoes_dg, s.decidido_por.nome if s.decidido_por else "",
        timezone.localtime(s.decidido_em).strftime("%d/%m/%Y %H:%M") if s.decidido_em else "",
        s.criado_por.nome)]


def _planilha(linhas) -> bytes:
    """A mesma exportação em XLSX: cabeçalho em negrito, congelado e com filtro."""
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    livro = Workbook()
    aba = livro.active if livro.active is not None else livro.create_sheet()
    aba.title = "Solicitações"
    aba.append(list(COLUNAS))
    for celula in aba[1]:
        celula.font = Font(bold=True)
    larguras = [len(c) for c in COLUNAS]
    for linha in linhas:
        aba.append(linha)
        larguras = [max(a, len(str(v))) for a, v in zip(larguras, linha, strict=False)]
    for i, largura in enumerate(larguras, start=1):
        aba.column_dimensions[get_column_letter(i)].width = min(max(10, largura + 2), 60)
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


@require_GET
def exportar(request: HttpRequest) -> HttpResponse:
    """A lista filtrada em CSV (";" e BOM) ou XLSX (`?formato=xlsx`), com as colunas da
    referência."""
    if not policies.pode_criar(request.user):
        raise PermissionDenied
    hoje = timezone.localdate()
    qs = (queries.filtrar(request.user, queries.ler_filtros(request.GET), hoje)
          .select_related("motorista", "decidido_por")
          .prefetch_related("servicos__servico", "equipes__equipe")
          .order_by("-data_solicitacao", "-criado_em"))
    linhas = (_linha_exportada(s) for s in qs)
    if request.GET.get("formato") == "xlsx":
        resposta = HttpResponse(_planilha(linhas), content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"))
        resposta["Content-Disposition"] = (
            f'attachment; filename="solicitacoes-{hoje:%Y-%m-%d}.xlsx"')
        return resposta
    resposta = HttpResponse(content_type="text/csv; charset=utf-8")
    resposta["Content-Disposition"] = f'attachment; filename="solicitacoes-{hoje:%Y-%m-%d}.csv"'
    resposta.write("\ufeff")
    escritor = csv.writer(resposta, delimiter=";", lineterminator="\r\n")
    escritor.writerow(COLUNAS)
    for linha in linhas:
        escritor.writerow(linha)
    return resposta


def _celula(valor):
    """Texto que começa como fórmula (=, +, -, @, tab, CR) vai com apóstrofo: a planilha
    não o executa. Números saem como estão."""
    if isinstance(valor, str) and valor[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{valor}"
    return valor


# ---------------------------------------------------------------- viagem (E4)
def _viagem(request: HttpRequest, s: Solicitacao) -> dict:
    integracao = ganchos.viagem()
    info = integracao.resumo(request.user, s) if integracao else None
    mostrar = bool(info) and (s.status in (dominio.DEFERIDA, dominio.ATENDIDA)
                              or bool(info and info["viagens"]))
    return {"viagem_info": info, "mostrar_viagem": mostrar}


@require_POST
def gerar_viagem(request: HttpRequest, pk: int) -> HttpResponse:
    s = _visivel(request, pk)
    integracao = ganchos.viagem()
    if integracao is None:
        raise Http404
    texto = request.POST.get("unidade") or ""
    unidade = int(texto) if re.fullmatch(r"[0-9]{1,9}", texto) else None
    try:
        viagem = integracao.gerar(request.user, s, unidade)
    except ValueError as exc:
        messages.error(request, str(exc))
    else:
        messages.success(request, f"Viagem #{viagem.pk} gerada; a equipe de viagens da "
                                  "unidade foi avisada.")
    return redirect(reverse("eventos:solicitacao", args=[pk]) + "#viagem")


# ---------------------------------------------------------------- folha
def _contexto(request: HttpRequest, form: FormularioSolicitacao, s: Solicitacao | None, *,
              reabrir: bool = False, erros_estrutura: list[str] | None = None) -> dict:
    hoje = timezone.localdate()
    # O que a DG escolheu antes de um erro no despacho volta uma vez só, nesta tela.
    despacho_escolhido = request.session.pop("despacho_escolhido", "")
    despacho_observacao = request.session.pop("despacho_observacao", "")
    if s is None:
        editavel = True
    else:
        editavel = (policies.pode_editar_dados(request.user, s)
                    or (reabrir and policies.pode_reabrir(request.user, s)))
    escolhidos = ({x.servico_id: x.observacao for x in s.servicos.all()} if s else {})
    equipes_atuais = ({x.equipe_id: x.quantidade_servidores for x in s.equipes.all()}
                      if s else {})
    if request.method == "POST" and erros_estrutura is not None:
        est, _e = estrutura_do_post(request.POST, s)
        escolhidos, equipes_atuais = est.servicos, est.equipes
    elif s is None and (tipo := form.initial.get("_tipo_modelo")) is not None:
        escolhidos = {x.pk: "" for x in tipo.servicos_sugeridos.all()}
        equipes_atuais = {x.equipe_id: x.quantidade for x in tipo.equipes_padrao.all()}
    servicos = Servico.objects.filter(ativo=True) | Servico.objects.filter(pk__in=escolhidos)
    equipes = Equipe.objects.filter(ativo=True) | Equipe.objects.filter(pk__in=equipes_atuais)
    ctx = {
        "form": form, "form_id": FORM_ID, "s": s, "editavel": editavel, "reabrir": reabrir,
        "servicos": [(x, x.pk in escolhidos) for x in servicos.distinct().order_by("nome")],
        "equipes": [(x, x.pk in equipes_atuais, equipes_atuais.get(x.pk))
                    for x in equipes.distinct().order_by("nome")],
        "erros_estrutura": erros_estrutura or [],
        "solicitantes": list(policies.solicitacoes_visiveis(request.user)
                             .exclude(solicitante_nome="").order_by("solicitante_nome")
                             .values_list("solicitante_nome", flat=True).distinct()[:300]),
        "tipos_com_modelo": TipoEvento.objects.filter(ativo=True).order_by("nome"),
    }
    if s is None:
        ctx["migalhas"] = _migalhas(("Nova solicitação", ""))
        return ctx
    ctx.update({
        "pedir_coffee": pedir_coffee.url(request.user, s),
        "titulo": queries.titulo(s), "etapas": dominio.etapas(s.status),
        "selo": dominio.selo_de_tempo(s.data_inicio_evento, s.data_fim_evento, s.status, hoje),
        "em_cima": dominio.em_cima_da_hora(s.data_solicitacao, s.data_inicio_evento),
        "erros_envio": (dominio.erros_do_envio(s.status, solicitacoes.dados_do_envio(s))
                        if s.status in dominio.EDITAVEIS else []),
        "pode_enviar": policies.pode_editar_dados(request.user, s),
        "pode_reabrir": policies.pode_reabrir(request.user, s),
        "pode_despachar": policies.pode_despachar(request.user) and s.status == dominio.AGUARDANDO,
        "pode_concluir": policies.pode_concluir(request.user, s),
        "erro_concluir": dominio.pode_concluir(s.status, s.data_fim_evento or
                                               s.data_inicio_evento, hoje),
        "pode_cancelar": policies.pode_cancelar(request.user, s),
        "pode_transferir": policies.pode_transferir(request.user, s),
        "pode_excluir": policies.pode_excluir(request.user, s),
        "pode_anexos": policies.pode_mexer_nos_anexos(request.user, s),
        "decisoes": dominio.DECISOES, "textos_despacho": TextoDespacho.objects.filter(ativo=True),
        "despacho_escolhido": despacho_escolhido, "despacho_observacao": despacho_observacao,
        "fila": (queries.aplicar_fila(policies.solicitacoes_visiveis(request.user), "despacho",
                                      request.user, hoje).count()
                 if policies.pode_despachar(request.user) else 0),
        "usuarios": get_user_model().objects.filter(is_active=True).exclude(
            pk=s.criado_por_id).order_by("nome") if policies.pode_transferir(request.user, s)
        else [],
        "historico": queries.historico(s),
        # Aviso, não bloqueio: motorista ou unidade móvel ocupados, pedido repetido.
        "avisos_agenda": conflitos.avisos_da_solicitacao(s),
        **_viagem(request, s),
        "migalhas": _migalhas((str(s), "")),
    })
    return ctx


def _erro(form, exc: dominio.RegraViolada) -> None:
    form.add_error(exc.campo if exc.campo in form.fields else None, str(exc))


@require_http_methods(["GET", "POST"])
def nova(request: HttpRequest) -> HttpResponse:
    if not policies.pode_criar(request.user):
        raise PermissionDenied
    if request.method == "POST":
        form = FormularioSolicitacao(request.POST)
        estrutura, erros = estrutura_do_post(request.POST)
        if form.is_valid() and not erros:
            try:
                nova_s = solicitacoes.criar(request.user, form.cleaned_data, estrutura)
            except dominio.RegraViolada as exc:
                _erro(form, exc)
            else:
                if request.POST.get("acao") != "enviar":
                    messages.success(request, f"Rascunho #{nova_s.pk} salvo com sucesso.")
                    return redirect("eventos:solicitacao", pk=nova_s.pk)
                try:
                    solicitacoes.enviar(request.user, nova_s.pk)
                except dominio.RegraViolada as exc:
                    messages.error(request, f"Rascunho #{nova_s.pk} salvo, mas não enviado: {exc}")
                else:
                    messages.success(request, f"Solicitação #{nova_s.pk} enviada com sucesso.")
                return redirect("eventos:solicitacao", pk=nova_s.pk)
        messages.error(request, "Corrija os campos destacados para continuar.")
        return render(request, "eventos/solicitacoes/folha.html",
                      _contexto(request, form, None, erros_estrutura=erros))
    inicial: dict = {"data_solicitacao": timezone.localdate()}
    for chave in ("inicio", "fim"):  # vindo da agenda: ?inicio=&fim=
        if (d := queries._data(request.GET.get(chave))) is not None:
            inicial["data_inicio_evento" if chave == "inicio" else "data_fim_evento"] = d
    tipo = TipoEvento.objects.filter(pk=request.GET.get("tipo") or 0, ativo=True).first() \
        if (request.GET.get("tipo") or "").isdigit() else None
    if tipo is not None:  # "Usar sugestão": o modelo do tipo
        inicial.update({"tipo_evento": tipo, "solicitante_nome": tipo.solicitante_padrao,
                        "solicitante_cargo_unidade": tipo.cargo_padrao,
                        "orgao_responsavel": tipo.orgao_padrao, "_tipo_modelo": tipo})
    form = FormularioSolicitacao(initial=inicial)
    return render(request, "eventos/solicitacoes/folha.html", _contexto(request, form, None))


@require_http_methods(["GET", "POST"])
def solicitacao(request: HttpRequest, pk: int) -> HttpResponse:
    s = _visivel(request, pk)
    reabrir = request.GET.get("reabrir") == "1" or request.POST.get("reabrir") == "1"
    if request.method == "POST":
        form = FormularioSolicitacao(request.POST, instance=s)
        estrutura, erros = estrutura_do_post(request.POST, s)
        if form.is_valid() and not erros:
            try:
                if reabrir:
                    _s, mudou = solicitacoes.reabrir_e_reenviar(request.user, s.pk,
                                                                form.cleaned_data, estrutura)
                    messages.success(request, (
                        f"Solicitação #{s.pk} alterada e reenviada para o despacho da DG."
                        if mudou else "Nada foi alterado: a solicitação continua como estava."))
                    return redirect("eventos:solicitacao", pk=s.pk)
                solicitacoes.salvar(request.user, s.pk, form.cleaned_data, estrutura)
                if request.POST.get("acao") == "enviar":
                    solicitacoes.enviar(request.user, s.pk)
                    messages.success(request, f"Solicitação #{s.pk} enviada com sucesso.")
                else:
                    messages.success(request, f"Solicitação #{s.pk} atualizada.")
                return redirect("eventos:solicitacao", pk=s.pk)
            except dominio.RegraViolada as exc:
                _erro(form, exc)
        messages.error(request, "Corrija os campos destacados para continuar.")
        s = _visivel(request, pk)
        return render(request, "eventos/solicitacoes/folha.html",
                      _contexto(request, form, s, reabrir=reabrir, erros_estrutura=erros))
    form = FormularioSolicitacao(instance=s)
    return render(request, "eventos/solicitacoes/folha.html",
                  _contexto(request, form, s, reabrir=reabrir))


@require_POST
def autosave(request: HttpRequest, pk: int) -> JsonResponse:
    s = _visivel(request, pk)
    if not policies.pode_editar_dados(request.user, s):
        return JsonResponse({"salvo": False, "mensagem": "Esta solicitação não pode ser alterada."})
    form = FormularioSolicitacao(request.POST, instance=s)
    estrutura, erros = estrutura_do_post(request.POST, s)
    if not form.is_valid() or erros:
        primeiro = erros[0] if erros else next(
            (f"{form.fields[n].label if n in form.fields else ''} — {e[0]}".strip(" —")
             for n, e in form.errors.items()), "")
        return JsonResponse({"salvo": False, "mensagem": f"Não salvo: {primeiro}"})
    try:
        salvo = solicitacoes.salvar(request.user, s.pk, form.cleaned_data, estrutura)
    except dominio.RegraViolada as exc:
        return JsonResponse({"salvo": False, "mensagem": f"Não salvo: {exc}"})
    from .forms import versao_de
    return JsonResponse({"salvo": True, "recarregar": False,
                         "em": timezone.localtime(salvo.atualizado_em).strftime("%H:%M"),
                         "campos": {"versao": versao_de(salvo)}})


# ---------------------------------------------------------------- ações
def _acao(request: HttpRequest, pk: int, faz, ok: str, ancora: str = "") -> HttpResponse:
    s = _visivel(request, pk)
    try:
        resultado = faz(s)
    except dominio.RegraViolada as exc:
        messages.error(request, str(exc))
    else:
        if ok:
            messages.success(request, ok.format(s=s, r=resultado))
    return redirect(reverse("eventos:solicitacao", args=[pk]) + ancora)


@require_POST
def enviar(request: HttpRequest, pk: int) -> HttpResponse:
    return _acao(request, pk, lambda s: solicitacoes.enviar(request.user, s.pk),
                 "Solicitação #{s.pk} enviada com sucesso.")


def _quantidades(request: HttpRequest) -> dict[int, int]:
    saida = {}
    for chave in request.POST:
        if chave.startswith("quantidade_dg_") and re.fullmatch(r"[0-9]{1,9}", chave[14:]):
            valor = quantidade(request.POST.get(chave))
            if valor is not None:
                saida[int(chave[14:])] = valor
    return saida


@require_POST
def despachar(request: HttpRequest, pk: int) -> HttpResponse:
    if not policies.pode_despachar(request.user):
        raise PermissionDenied
    s = _visivel(request, pk)
    if request.POST.get("acao_despacho") == "salvar_ajustes":
        try:
            mudancas = solicitacoes.ajustar_servidores(request.user, s.pk, _quantidades(request))
        except dominio.RegraViolada as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, (f"Ajustes salvos para a solicitação #{s.pk}: "
                                       + "; ".join(mudancas) + ".") if mudancas
                             else "Nenhuma alteração nas quantidades.")
        return redirect(reverse("eventos:solicitacao", args=[pk]) + "#despacho")
    decisao = request.POST.get("decisao") or ""
    try:
        feita = solicitacoes.despachar(request.user, s.pk, decisao,
                                       request.POST.get("observacao") or "", _quantidades(request))
    except dominio.RegraViolada as exc:
        messages.error(request, str(exc))
        request.session["despacho_escolhido"] = decisao
        request.session["despacho_observacao"] = (request.POST.get("observacao") or "")[:2000]
        return redirect(reverse("eventos:solicitacao", args=[pk]) + "#despacho")
    messages.success(request, f"Solicitação #{feita.pk} enviada para correção." if decisao ==
                     "devolver" else f"Decisão registrada para a solicitação #{feita.pk}.")
    if request.POST.get("seguir") == "proxima":
        proxima = solicitacoes.proxima_da_fila(request.user, depois_de=feita.pk)
        if proxima is None:
            messages.info(request, "Não há mais solicitações aguardando despacho.")
            return redirect(reverse("eventos:solicitacoes") + "?fila=despacho")
        return redirect(reverse("eventos:solicitacao", args=[proxima.pk]) + "#despacho")
    return redirect("eventos:solicitacao", pk=feita.pk)


@require_POST
def concluir(request: HttpRequest, pk: int) -> HttpResponse:
    return _acao(request, pk, lambda s: solicitacoes.concluir(request.user, s.pk),
                 "Atendimento da solicitação #{s.pk} confirmado.")


@require_POST
def cancelar(request: HttpRequest, pk: int) -> HttpResponse:
    return _acao(request, pk, lambda s: solicitacoes.cancelar(
        request.user, s.pk, request.POST.get("motivo_cancelamento") or ""),
        "Evento da solicitação #{s.pk} registrado como cancelado.", "#encerramento")


@require_POST
def transferir(request: HttpRequest, pk: int) -> HttpResponse:
    pk_novo = request.POST.get("responsavel") or ""
    novo = (get_user_model().objects.filter(pk=pk_novo).first()
            if re.fullmatch(r"[0-9]{1,9}", pk_novo) else None)
    return _acao(request, pk, lambda s: solicitacoes.transferir(
        request.user, s.pk, novo, request.POST.get("motivo_transferencia") or ""),
        "Solicitação #{s.pk} transferida para {r.criado_por.nome}.", "#responsavel")


@require_POST
def duplicar(request: HttpRequest, pk: int) -> HttpResponse:
    s = _visivel(request, pk)
    copia = solicitacoes.duplicar(request.user, s.pk)
    messages.success(request, f"Rascunho #{copia.pk} criado a partir da solicitação #{s.pk}. "
                              "Informe as datas do evento e envie.")
    return redirect("eventos:solicitacao", pk=copia.pk)


@require_POST
def excluir(request: HttpRequest, pk: int) -> HttpResponse:
    s = _visivel(request, pk)
    solicitacoes.excluir(request.user, s.pk)
    messages.success(request, f"Rascunho #{pk} excluído.")
    return redirect("eventos:solicitacoes")


@require_POST
def anexar(request: HttpRequest, pk: int) -> HttpResponse:
    s = _visivel(request, pk)
    arquivos = request.FILES.getlist("arquivo")
    if not arquivos:
        messages.error(request, "Selecione um arquivo para anexar.")
    for arquivo in arquivos:
        try:
            solicitacoes.anexar(request.user, s.pk, arquivo)
        except dominio.RegraViolada as exc:
            messages.error(request, str(exc))
    return redirect(reverse("eventos:solicitacao", args=[pk]) + "#anexos")


@require_POST
def remover_anexo(request: HttpRequest, pk: int, anexo_pk: int) -> HttpResponse:
    _visivel(request, pk)
    if not AnexoSolicitacao.objects.filter(pk=anexo_pk, solicitacao_id=pk).exists():
        raise Http404
    nome = solicitacoes.remover_anexo(request.user, anexo_pk)
    messages.success(request, f"Anexo removido: {nome}.")
    return redirect(reverse("eventos:solicitacao", args=[pk]) + "#anexos")


@require_GET
def abrir_anexo(request: HttpRequest, pk: int, anexo_pk: int) -> FileResponse:
    _visivel(request, pk)
    anexo = get_object_or_404(AnexoSolicitacao, pk=anexo_pk, solicitacao_id=pk)
    resposta = FileResponse(anexo.arquivo.open("rb"), as_attachment=True,
                            filename=anexo.nome_original)
    resposta["Cache-Control"] = "no-store"
    resposta["X-Content-Type-Options"] = "nosniff"
    return resposta
