"""Solicitação de evento de ponta a ponta: rascunho pela tela (serviços e equipes),
autosave, envio (o que falta), visibilidade, despacho da DG (decisões, observação,
devolução, ajuste de servidores, "abrir a próxima"), reenvio depois do despacho,
atendida só depois do evento, cancelar, transferir, duplicar, excluir, anexos, lista,
exportação, agenda e conflitos."""

from __future__ import annotations

import csv
import io
from datetime import timedelta
from pathlib import Path

import pytest
from django.contrib.auth.models import Group
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client
from django.urls import reverse
from django.utils import timezone

from gestao.cadastros.models import Municipio
from gestao.eventos import conflitos, solicitacoes
from gestao.eventos.models import (
    Equipe,
    Movimento,
    OrgaoResponsavel,
    Servico,
    Solicitacao,
    TipoEvento,
    UnidadeMovel,
)
from gestao.identidade.models import Usuario
from gestao.plataforma import agenda
from gestao.plataforma.models import Notificacao

pytestmark = pytest.mark.django_db


def _usuario(login: str, *papeis: str) -> Usuario:
    u = Usuario.objects.create_user(login, f"{login}@teste.invalid", None, nome=login.title())
    for p in papeis:
        u.groups.add(Group.objects.get(name=p))
    return u


@pytest.fixture
def ana():
    return _usuario("ana")


@pytest.fixture
def dg():
    return _usuario("dg", "GESTOR_DG")


@pytest.fixture
def curitiba():
    return Municipio.objects.get_or_create(codigo_ibge="4106902",
                                           defaults={"nome": "Curitiba", "uf": "PR"})[0]


def _completa(**extra) -> dict:
    hoje = timezone.localdate()
    return {"data_solicitacao": hoje, "data_inicio_evento": hoje + timedelta(days=20),
            "data_fim_evento": hoje + timedelta(days=21),
            "municipio": Municipio.objects.get_or_create(
                codigo_ibge="4106902", defaults={"nome": "Curitiba", "uf": "PR"})[0],
            "tipo_evento": TipoEvento.objects.get(nome="Feira"),
            "solicitante_nome": "Escola X", "solicitante_cargo_unidade": "Direção",
            "orgao_responsavel": OrgaoResponsavel.objects.first(), "local_evento": "Ginásio",
            "tipo_operacao": "diaria", **extra}


def _estrutura(qtd: int | None = 3) -> solicitacoes.Estrutura:
    return solicitacoes.Estrutura(
        servicos={Servico.objects.get(nome="Emissão de CIN").pk: ""},
        equipes={Equipe.objects.get(nome="Alfa").pk: qtd})


def test_rascunho_pela_tela_e_autosave(ana, curitiba):
    c = Client()
    c.force_login(ana)
    cin = Servico.objects.get(nome="Emissão de CIN")
    alfa = Equipe.objects.get(nome="Alfa")
    hoje = timezone.localdate()
    r = c.post(reverse("eventos:nova"), {
        "data_solicitacao": f"{hoje:%d/%m/%Y}", "municipio": "Curitiba/PR",
        "solicitante_nome": "  Escola   X ", "tipo_operacao": "diaria", "acao": "rascunho",
        f"servico_{cin.pk}": "1", f"equipe_{alfa.pk}": "1", f"quantidade_equipe_{alfa.pk}": "4"})
    s = Solicitacao.objects.get()
    assert r.status_code == 302 and s.status == "rascunho" and s.criado_por == ana
    assert s.solicitante_nome == "Escola X" and s.municipio == curitiba
    assert s.quantidade_servidores == 4 and s.servicos.count() == 1
    assert Movimento.objects.filter(solicitacao=s, acao="criacao").exists()
    from gestao.eventos.forms import versao_de
    r = c.post(reverse("eventos:autosave", args=[s.pk]), {
        "versao": versao_de(s), "data_solicitacao": f"{hoje:%d/%m/%Y}",
        "solicitante_nome": "Escola Y", "tipo_operacao": "diaria", "municipio": "Curitiba/PR",
        f"equipe_{alfa.pk}": "1", f"quantidade_equipe_{alfa.pk}": "zero"})
    assert r.json()["salvo"] is False and "quantidade válida" in r.json()["mensagem"]
    r = c.post(reverse("eventos:autosave", args=[s.pk]), {
        "versao": "123", "data_solicitacao": f"{hoje:%d/%m/%Y}", "tipo_operacao": "diaria"})
    assert r.json()["salvo"] is False and "outra pessoa" in r.json()["mensagem"]


def test_envio_exige_o_necessario_e_avisa_a_dg(ana, dg, django_capture_on_commit_callbacks):
    s = solicitacoes.criar(ana, {"solicitante_nome": "X"})
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="Preencha os campos obrigatórios"):
        solicitacoes.enviar(ana, s.pk)
    solicitacoes.salvar(ana, s.pk, _completa(), _estrutura(None))
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="quantidade de servidores de Alfa"):
        solicitacoes.enviar(ana, s.pk)
    solicitacoes.salvar(ana, s.pk, {}, _estrutura(3))
    with pytest.raises(PermissionDenied):
        solicitacoes.enviar(dg, s.pk)  # só o responsável envia
    with django_capture_on_commit_callbacks(execute=True):
        solicitacoes.enviar(ana, s.pk)
    s.refresh_from_db()
    assert s.status == "aguardando_despacho"
    assert Notificacao.objects.filter(usuario=dg, titulo=f"Solicitação #{s.pk} aguardando despacho")
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="Apenas rascunhos"):
        solicitacoes.enviar(ana, s.pk)


def test_visibilidade(ana, dg):
    bia = _usuario("bia")
    s = solicitacoes.criar(ana, _completa(), _estrutura())
    c = Client()
    c.force_login(bia)
    assert c.get(reverse("eventos:solicitacao", args=[s.pk])).status_code == 404
    assert "Escola X" not in c.get(reverse("eventos:solicitacoes")).content.decode()
    c.force_login(dg)
    assert c.get(reverse("eventos:solicitacao", args=[s.pk])).status_code == 200


def test_despacho_devolucao_reenvio_e_proxima(ana, dg, django_capture_on_commit_callbacks):
    s = solicitacoes.criar(ana, _completa(), _estrutura())
    solicitacoes.enviar(ana, s.pk)
    with pytest.raises(PermissionDenied):
        solicitacoes.despachar(ana, s.pk, "atender")
    admin = _usuario("adm", "ADMINISTRADOR")
    with pytest.raises(PermissionDenied):
        solicitacoes.despachar(admin, s.pk, "atender")  # quem administra não despacha
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="observação é obrigatória"):
        solicitacoes.despachar(dg, s.pk, "nao_atender", "")
    c = Client()
    c.force_login(dg)
    alfa = Equipe.objects.get(nome="Alfa")
    with django_capture_on_commit_callbacks(execute=True):
        r = c.post(reverse("eventos:despachar", args=[s.pk]), {
            "decisao": "devolver", "observacao": "Falta o endereço",
            f"quantidade_dg_{alfa.pk}": "5"})
    assert r.status_code == 302
    s.refresh_from_db()
    assert s.status == "devolvida" and s.quantidade_servidores == 5
    assert Notificacao.objects.filter(usuario=ana, titulo__contains="enviada para correção")
    solicitacoes.enviar(ana, s.pk)
    outra = solicitacoes.criar(ana, _completa(), _estrutura())
    solicitacoes.enviar(ana, outra.pk)
    r = c.post(reverse("eventos:despachar", args=[s.pk]), {"decisao": "atender",
                                                          "seguir": "proxima"})
    assert r["Location"].startswith(reverse("eventos:solicitacao", args=[outra.pk]))
    s.refresh_from_db()
    assert s.status == "deferida" and s.decisao_dg == "atender" and s.decidido_por == dg
    # Alterar depois do despacho: volta a aguardar e a decisão sai.
    _s, mudou = solicitacoes.reabrir_e_reenviar(ana, s.pk, {"local_evento": "Praça"})
    s.refresh_from_db()
    assert mudou and s.status == "aguardando_despacho" and s.decisao_dg == "pendente"
    assert Movimento.objects.filter(solicitacao=s, acao="reenvio",
                                    observacao__contains="aguarda novo despacho").exists()
    _s, mudou = solicitacoes.reabrir_e_reenviar(ana, s.pk, {"local_evento": "Praça"})
    assert not mudou


def test_concluir_cancelar_transferir_duplicar_excluir(ana, dg,
                                                      django_capture_on_commit_callbacks):
    hoje = timezone.localdate()
    s = solicitacoes.criar(ana, _completa(), _estrutura())
    solicitacoes.enviar(ana, s.pk)
    solicitacoes.despachar(dg, s.pk, "atender")
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="depois que o evento terminar"):
        solicitacoes.concluir(ana, s.pk)
    solicitacoes.concluir(ana, s.pk, hoje=hoje + timedelta(days=30))
    s.refresh_from_db()
    assert s.status == "atendida"
    t = solicitacoes.criar(ana, _completa(), _estrutura())
    solicitacoes.enviar(ana, t.pk)
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="motivo do cancelamento"):
        solicitacoes.cancelar(ana, t.pk, "")
    solicitacoes.cancelar(dg, t.pk, "Chuva")
    t.refresh_from_db()
    assert t.status == "cancelada"
    bia = _usuario("bia")
    u = solicitacoes.criar(ana, _completa(), _estrutura())
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="já é a responsável"):
        solicitacoes.transferir(ana, u.pk, ana)
    with django_capture_on_commit_callbacks(execute=True):
        solicitacoes.transferir(ana, u.pk, bia, "Férias")
    u.refresh_from_db()
    assert u.criado_por == bia and Notificacao.objects.filter(usuario=bia).exists()
    copia = solicitacoes.duplicar(dg, s.pk)
    assert copia.status == "rascunho" and copia.criado_por == dg
    assert copia.data_inicio_evento is None and copia.servicos.count() == 1
    with pytest.raises(PermissionDenied):
        solicitacoes.excluir(ana, s.pk)  # não é rascunho
    solicitacoes.excluir(dg, copia.pk)
    assert not Solicitacao.objects.filter(pk=copia.pk).exists()


def test_anexos(ana, settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path
    s = solicitacoes.criar(ana, _completa(), _estrutura())
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="tipo de arquivo não permitido"):
        solicitacoes.anexar(ana, s.pk, SimpleUploadedFile("x.exe", b"MZ"))
    with pytest.raises(solicitacoes.SolicitacaoInvalida, match="não corresponde ao tipo"):
        solicitacoes.anexar(ana, s.pk, SimpleUploadedFile("x.pdf", b"nao e pdf"))
    a = solicitacoes.anexar(ana, s.pk, SimpleUploadedFile("oficio.pdf", b"%PDF-1.4 teste"))
    c = Client()
    c.force_login(ana)
    r = c.get(reverse("eventos:abrir_anexo", args=[s.pk, a.pk]))
    assert r.status_code == 200 and b"".join(r.streaming_content).startswith(b"%PDF")
    solicitacoes.remover_anexo(ana, a.pk)
    assert not s.anexos.exists()


def test_lista_filas_busca_e_exportacao(ana, dg):
    a = solicitacoes.criar(ana, _completa(solicitante_nome="Colégio Azul"), _estrutura())
    b = solicitacoes.criar(ana, _completa(solicitante_nome="Igreja Verde"), _estrutura())
    solicitacoes.enviar(ana, b.pk)
    c = Client()
    c.force_login(dg)
    html = c.get(reverse("eventos:solicitacoes")).content.decode()
    assert 'Aguardando despacho <span class="aba__contagem">1</span>' in html
    html = c.get(reverse("eventos:solicitacoes"), {"q": "azul"}).content.decode()
    assert "Colégio Azul" in html and "Igreja Verde" not in html
    html = c.get(reverse("eventos:solicitacoes"), {"q": f"#{b.pk}"}).content.decode()
    assert f"Pedido</span><span class=\"placa__numero\">#{b.pk}" in html
    c.force_login(ana)
    html = c.get(reverse("eventos:solicitacoes"), {"fila": "rascunhos"}).content.decode()
    assert f"#{a.pk}" in html and "Aguardando despacho <span" not in html  # fila só da DG
    r = c.get(reverse("eventos:exportar"))
    linhas = list(csv.reader(io.StringIO(r.content.decode().lstrip("﻿")), delimiter=";"))
    assert linhas[0][0] == "Nº" and len(linhas) == 3
    # A mesma exportação em XLSX (E5): mesmas colunas e linhas, fórmula neutralizada.
    from openpyxl import load_workbook
    r = c.get(reverse("eventos:exportar"), {"formato": "xlsx"})
    assert r["Content-Type"].startswith("application/vnd.openxmlformats")
    aba = load_workbook(io.BytesIO(r.content)).active
    assert aba["A1"].value == "Nº" and aba.max_row == 3
    assert "formato=xlsx" in c.get(reverse("eventos:solicitacoes")).content.decode()


def test_agenda_e_conflitos(ana, dg):
    um = UnidadeMovel.objects.create(nome="Caminhão 1")
    a = solicitacoes.criar(ana, _completa(unidade_movel=True, unidade_movel_designada=um),
                           _estrutura())
    solicitacoes.enviar(ana, a.pk)
    b = solicitacoes.criar(ana, _completa(unidade_movel=True, unidade_movel_designada=um),
                           _estrutura())
    avisos = conflitos.avisos_da_solicitacao(b)
    assert {x.tipo for x in avisos} == {"unidade_movel", "pedido"}
    assert all(x.chave == ("solicitacao", a.pk) for x in avisos)
    inicio = a.data_inicio_evento
    itens = agenda.compromissos_de(dg, inicio, inicio, ["solicitacao"])
    assert [i.chave for i in itens] == [f"solicitacao-{a.pk}"]  # rascunho não entra


def test_entradas_hostis_nao_quebram_nem_viram_formula(ana, dg):
    """Revisão de segurança do E2: fórmula em qualquer coluna do CSV, algarismos que não
    são ASCII, quantidade enorme, serviço inativo pelo POST e anexo que só some depois do
    commit."""
    from gestao.eventos.forms import estrutura_do_post, quantidade

    s = solicitacoes.criar(ana, _completa(bairro='=HYPERLINK("https://x.invalid")',
                                          cep="\t=1+1"), _estrutura())
    c = Client()
    c.force_login(ana)
    r = c.get(reverse("eventos:exportar"))
    linha = next(row for row in csv.reader(
        io.StringIO(r.content.decode().lstrip("\ufeff")), delimiter=";") if row[0] == str(s.pk))
    assert linha[9].startswith("'=") and linha[10].startswith("'\t")
    for q in ("#²", "²"):
        assert c.get(reverse("eventos:solicitacoes"), {"q": q, "municipio": "²"}).status_code == 200
    assert quantidade("²") is None and quantidade("99999999999") is None
    assert quantidade("12") == 12
    velho = Servico.objects.create(nome="Serviço antigo", ativo=False)
    alfa = Equipe.objects.get(nome="Alfa")
    estrutura, erros = estrutura_do_post({f"servico_{velho.pk}": "1", f"equipe_{alfa.pk}": "1",
                                          f"quantidade_equipe_{alfa.pk}": "99999999999"})
    assert velho.pk not in estrutura.servicos and erros  # inativo fora; quantidade recusada
    r = Client()
    r.force_login(dg)
    solicitacoes.enviar(ana, s.pk)
    resposta = r.post(reverse("eventos:despachar", args=[s.pk]), {
        "decisao": "atender", f"quantidade_dg_{alfa.pk}": "99999999999"})
    assert resposta.status_code == 302  # quantidade fora do limite é ignorada, não 500


def test_excluir_so_apaga_os_arquivos_depois_do_commit(ana, settings, tmp_path,
                                                         django_capture_on_commit_callbacks):
    settings.MEDIA_ROOT = tmp_path
    s = solicitacoes.criar(ana, {"solicitante_nome": "X"})
    a = solicitacoes.anexar(ana, s.pk, SimpleUploadedFile("o.pdf", b"%PDF-1.4 teste"))
    caminho = Path(a.arquivo.path)
    with django_capture_on_commit_callbacks(execute=False) as pendentes:
        solicitacoes.excluir(ana, s.pk)
    assert caminho.exists() and pendentes  # ainda não: espera o commit
    for f in pendentes:
        f()
    assert not caminho.exists()


def test_fila_do_despacho_pelo_evento_mais_proximo_e_escolha_que_volta_uma_vez(ana, dg):
    """Revisão de UX do E2: a fila do despacho na ordem do "abrir a próxima"; depois de um
    erro no despacho a decisão e a observação voltam na tela, uma vez só."""
    hoje = timezone.localdate()
    longe = solicitacoes.criar(ana, _completa(data_inicio_evento=hoje + timedelta(days=40),
                                              data_fim_evento=hoje + timedelta(days=40)),
                               _estrutura())
    perto = solicitacoes.criar(ana, _completa(data_inicio_evento=hoje + timedelta(days=5),
                                              data_fim_evento=hoje + timedelta(days=5)),
                               _estrutura())
    for s in (longe, perto):
        solicitacoes.enviar(ana, s.pk)
    c = Client()
    c.force_login(dg)
    html = c.get(reverse("eventos:solicitacoes"), {"fila": "despacho"}).content.decode()
    assert html.index(f">#{perto.pk}<") < html.index(f">#{longe.pk}<")
    assert "Despachar" in html
    r = c.post(reverse("eventos:despachar", args=[perto.pk]),
               {"decisao": "nao_atender", "observacao": ""})  # falta a observação
    html = c.get(r["Location"]).content.decode()
    assert 'value="nao_atender" checked' in html
    html = c.get(reverse("eventos:solicitacao", args=[perto.pk])).content.decode()
    assert 'value="nao_atender" checked' not in html  # só uma vez
    from gestao.eventos import dominio
    assert dominio.selo_de_tempo(hoje - timedelta(days=9), None, "cancelada", hoje) is None


def test_rodape_padrao_e_cancelar_pelo_dialogo(ana, dg):
    """Rodapé igual em toda folha (06/10): Finalizar e Ações (Duplicar, Cancelar, Excluir);
    o cancelamento pela janela de motivo manda `motivo`."""
    s = solicitacoes.criar(ana, _completa(), _estrutura())
    c = Client()
    c.force_login(ana)
    html = c.get(reverse("eventos:solicitacao", args=[s.pk])).content.decode()
    assert 'form="finalizar-evento"' in html and "Duplicar solicitação" in html
    assert "Excluir solicitação" in html  # rascunho
    solicitacoes.enviar(ana, s.pk)  # aguardando despacho: a responsável pode cancelar
    html = c.get(reverse("eventos:solicitacao", args=[s.pk])).content.decode()
    assert "Cancelar solicitação" in html and "Excluir solicitação" not in html
    c.post(reverse("eventos:cancelar", args=[s.pk]), {"motivo": "Evento adiado"})
    s.refresh_from_db()
    assert s.status == Solicitacao.Status.CANCELADA
